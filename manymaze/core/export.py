"""Export results (CSV / tab-separated / Excel / HTML report), per-test raw data and whole experiments as XML."""

from __future__ import annotations

import base64
import csv
import re
import datetime as _dt
import html
import logging
import math
from pathlib import Path
from xml.sax.saxutils import escape as _xesc, quoteattr as _qa

import numpy as np

from .. import __version__
from .atomicfile import atomic_write, write_text_atomic
from .apparatus import CALIBRATION_KEY, ENTRY_RULES as ENTRY_RULE_TEXT, POSITION_KEY, position_args
from .project import INACTIVE_STATUSES, INFO_COLUMNS, Project, result_columns
from .stats import is_number

XML_FORMAT_VERSION = 1
log = logging.getLogger(__name__)


def value_text(v) -> str:
    """Full-precision text of a value (exports, clipboard): whole numbers as integers, other numbers with the shortest
    text that reads back exactly (repr), blank for missing and non-finite values."""
    if v is None:
        return ""
    if isinstance(v, (float, np.floating)):
        if not math.isfinite(v):
            return ""
        f = float(v)
        return str(int(f)) if f.is_integer() and abs(f) < 1e15 else repr(f)
    if isinstance(v, np.integer):
        return str(int(v))
    return str(v)


def display_text(v) -> str:
    """A value as shown in tables: numbers with 3 decimals (blank if not finite)."""
    if v is None:
        return ""
    if is_number(v):
        if isinstance(v, (int, np.integer)):
            return str(int(v))
        return f"{float(v):.3f}" if math.isfinite(v) else ""
    return str(v)


def _csv_text(v) -> str:
    """value_text, with text that looks like a formula prefixed by ' so a spreadsheet keeps it as text (no formula
    injection; as the xlsx writer). Numbers, negative ones included, and plain labels such as "+/+" or "-ctrl" are
    left alone: a leading + or - counts only when the text also has formula syntax (parentheses, ! : \\ |)."""
    s = value_text(v)
    if not isinstance(v, str) or not s:
        return s
    if s[0] in ("=", "@") or (s[0] in "+-" and re.search(r"[(!:\\|]", s)):
        return "'" + s
    return s


def write_csv(rows: list[dict], path, columns: list[str] | None = None, delimiter=","):
    cols = columns or result_columns(rows)
    with atomic_write(path, newline="") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow([_csv_text(c) for c in cols])
        for r in rows:
            w.writerow([_csv_text(r.get(c)) for c in cols])


def write_tsv(rows: list[dict], path, columns: list[str] | None = None):
    """Tab-separated text (opens directly in Excel / Prism / R)."""
    write_csv(rows, path, columns, delimiter="\t")


def _cell_number(v):
    """v as a finite int / float for a numeric cell, else None (text, bool, missing, NaN / inf)."""
    if isinstance(v, (bool, np.bool_)) or not is_number(v):
        return None
    f = float(v)
    return (int(f) if isinstance(v, (int, np.integer)) else f) if math.isfinite(f) else None


def _sylk_text(s: str) -> str:
    # one record per line; ';' separates fields so a literal ';' is doubled
    return re.sub(r"[\x00-\x1f]", " ", s).replace(";", ";;")


def write_sylk(rows: list[dict], path, columns: list[str] | None = None):
    """SYLK (Symbolic Link, .slk): the plain-text spreadsheet format of Multiplan / Excel. Numbers are stored as
    numbers, everything else as text; the column headings form the first row (bold). Written in Windows-1252, the
    encoding Excel expects for SYLK."""
    cols = columns or result_columns(rows)
    out = ["ID;PmANY-MAZE;N;E", f"B;Y{len(rows) + 1};X{max(1, len(cols))};D0 0 {len(rows)} {max(0, len(cols) - 1)}",
           "F;SD;R1"]  # SD R1: the heading row is bold
    for x, c in enumerate(cols, 1):
        out.append(f'C;Y1;X{x};K"{_sylk_text(str(c))}"')
    for y, r in enumerate(rows, 2):
        first = True
        for x, c in enumerate(cols, 1):
            v = r.get(c)
            n = _cell_number(v)
            if n is not None:
                val = value_text(n)
            elif isinstance(v, (bool, np.bool_)):
                val = "TRUE" if v else "FALSE"
            elif v is None or value_text(v) == "":
                continue
            else:
                val = f'"{_sylk_text(value_text(v))}"'
            out.append(f"C;{f'Y{y};' if first else ''}X{x};K{val}")
            first = False
    out.append("E")
    with atomic_write(path, "wb") as f:
        f.write(("\r\n".join(out) + "\r\n").encode("cp1252", errors="replace"))


def read_sylk(path) -> list[list]:
    """Cells of a SYLK file as a list of rows (numbers as float, text as str, empty cells None) — enough to read
    back what write_sylk writes (and simple SYLK files of other programs)."""
    grid: dict[tuple[int, int], object] = {}
    y = x = 1
    text = Path(path).read_bytes().decode("cp1252", errors="replace")
    for line in text.splitlines():
        fields = re.split(r"(?<!;);(?!;)", line)  # split on single ';', not ';;'
        if not fields or fields[0] != "C":
            continue
        val = None
        for f in fields[1:]:
            if f[:1] == "Y":
                y = int(f[1:])
            elif f[:1] == "X":
                x = int(f[1:])
            elif f[:1] == "K":
                k = f[1:].replace(";;", ";")
                if k.startswith('"'):
                    val = k[1:-1] if k.endswith('"') else k[1:]
                elif k in ("TRUE", "FALSE"):
                    val = k == "TRUE"
                else:
                    val = float(k)
        grid[(y, x)] = val
    if not grid:
        return []
    ny, nx = max(k[0] for k in grid), max(k[1] for k in grid)
    return [[grid.get((j, i)) for i in range(1, nx + 1)] for j in range(1, ny + 1)]


DBF_TEXT_MAX = 254
DBF_MAX_FIELDS = 255  # dBase IV (dBase III: 128)
_DBF_NUM_MAX = 19  # dBase III numeric field width limit (dBase IV: 20)


def dbf_field_names(columns: list[str]) -> list[str]:
    """dBase field names for column headings: letters, digits and '_', starting with a letter, at most 10 characters,
    upper case and unique (e.g. 'Total distance (m)' -> 'TOTAL_DIST', then 'TOTAL_DI_2')."""
    out: list[str] = []
    for c in columns:
        base = re.sub(r"_+", "_", re.sub(r"[^A-Za-z0-9]", "_", str(c))).strip("_").upper() or "FIELD"
        if not base[0].isalpha():
            base = "F" + base
        name, k = base[:10], 1
        while name in out:
            k += 1
            suffix = f"_{k}"
            name = base[:10 - len(suffix)] + suffix
        out.append(name)
    return out


def _dbf_field(values: list) -> tuple[str, int, int, list[bytes]]:
    """(type, length, decimals, encoded values) of one column: N (numbers), L (true / false) or C (text)."""
    present = [v for v in values if v is not None and value_text(v) != ""]
    if present and all(isinstance(v, (bool, np.bool_)) for v in present):
        return "L", 1, 0, [b"?" if v is None or value_text(v) == "" else (b"T" if v else b"F") for v in values]
    nums = [_cell_number(v) for v in values]
    if present and all(n is not None for v, n in zip(values, nums) if v is not None and value_text(v) != ""):
        texts = [value_text(n) for n in nums if n is not None]
        if not any("e" in t for t in texts):
            ints = max(len(t.split(".")[0]) for t in texts)
            dec = max((len(t.split(".")[1]) for t in texts if "." in t), default=0)
            dec = min(dec, max(0, _DBF_NUM_MAX - ints - 1), 15)
            width = ints + (dec + 1 if dec else 0)
            if width <= _DBF_NUM_MAX:
                enc = [(f"{n:.{dec}f}" if dec else str(int(round(n)))).rjust(width).encode("ascii")
                       if n is not None else b" " * width for n in nums]
                return "N", width, dec, enc
    enc = [("" if v is None else re.sub(r"[\r\n\t]", " ", value_text(v))).encode("cp1252", errors="replace")
           [:DBF_TEXT_MAX] for v in values]
    width = max([1] + [len(b) for b in enc])
    return "C", width, 0, [b.ljust(width) for b in enc]


def write_dbf(rows: list[dict], path, columns: list[str] | None = None, date: _dt.date | None = None):
    """dBase III table (.dbf, also read by dBase IV, Excel, LibreOffice, SPSS, R `foreign`, GIS software). Columns of
    numbers become numeric (N) fields, true / false columns logical (L), others text (C, at most 254 characters,
    Windows-1252). Field names are limited to 10 characters (see dbf_field_names); missing numbers are blank. At most
    DBF_MAX_FIELDS columns (dBase IV's limit; dBase III's is 128)."""
    cols = columns or result_columns(rows)
    names = dbf_field_names(cols)
    fields = [_dbf_field([r.get(c) for r in rows]) for c in cols]
    if len(cols) > DBF_MAX_FIELDS:
        raise ValueError(f"A dBase table holds at most {DBF_MAX_FIELDS} fields and this table has {len(cols)} "
                         "columns: show fewer measures (Select data) or save as CSV / Excel")
    record_len = 1 + sum(f[1] for f in fields)
    if record_len > 65535:
        raise ValueError("The rows are too wide for a dBase table")
    header_len = 32 + 32 * len(fields) + 1
    d = date or _dt.date.today()
    head = bytearray(32)
    head[0] = 0x03  # dBase III without memo
    head[1:4] = bytes((d.year - 1900, d.month, d.day))
    head[4:8] = len(rows).to_bytes(4, "little")
    head[8:10] = header_len.to_bytes(2, "little")
    head[10:12] = record_len.to_bytes(2, "little")
    head[29] = 0x57  # language driver: Windows ANSI (code page 1252)
    out = bytearray(head)
    for name, (typ, length, dec, _) in zip(names, fields):
        fd = bytearray(32)
        fd[0:len(name)] = name.encode("ascii")
        fd[11] = ord(typ)
        fd[16] = length
        fd[17] = dec
        out += fd
    out += b"\x0d"
    for i in range(len(rows)):
        out += b" " + b"".join(f[3][i] for f in fields)  # ' ': record not deleted
    out += b"\x1a"
    with atomic_write(path, "wb") as f:
        f.write(bytes(out))


def read_dbf(path) -> tuple[list[str], list[list]]:
    """Field names and records of a dBase III / IV table (.dbf): numbers as float (None if blank), logicals as
    bool (None if unknown), dates as 'YYYY-MM-DD', text as str; deleted records are skipped. Memo fields are not
    read (blank)."""
    raw = Path(path).read_bytes()
    if len(raw) < 33:
        raise ValueError(f"{Path(path).name} is not a dBase table")
    n, hlen, rlen = int.from_bytes(raw[4:8], "little"), int.from_bytes(raw[8:10], "little"), \
        int.from_bytes(raw[10:12], "little")
    enc = {0x57: "cp1252", 0x03: "cp1252", 0x01: "cp437", 0x02: "cp850", 0x64: "cp852", 0x65: "cp866",
           0xC8: "cp1250", 0xC9: "cp1251"}.get(raw[29], "cp1252")
    fields, pos = [], 32
    while pos + 32 <= hlen and raw[pos] != 0x0D:
        fd = raw[pos:pos + 32]
        fields.append((fd[:11].split(b"\0")[0].decode("ascii", "replace"), chr(fd[11]), fd[16]))
        pos += 32
    rows = []
    for i in range(n):
        rec = raw[hlen + i * rlen:hlen + (i + 1) * rlen]
        if len(rec) < rlen or rec[:1] == b"*":
            continue
        off, vals = 1, []
        for _, typ, length in fields:
            b = rec[off:off + length]
            off += length
            txt = b.decode(enc, "replace").strip()
            if typ in "NF":
                try:
                    vals.append(float(txt) if txt else None)
                except ValueError:
                    vals.append(None)
            elif typ == "L":
                vals.append(True if txt in ("T", "t", "Y", "y") else False if txt in ("F", "f", "N", "n") else None)
            elif typ == "D":
                vals.append(f"{txt[:4]}-{txt[4:6]}-{txt[6:8]}" if len(txt) == 8 and txt.isdigit() else "")
            elif typ == "M":
                vals.append("")
            else:
                vals.append(txt)
        rows.append(vals)
    return [f[0] for f in fields], rows


TABLE_SUFFIXES = (".csv", ".tsv", ".txt", ".xlsx", ".slk", ".dbf")


_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def write_xlsx(sheets: dict[str, list[dict]], path, columns: dict[str, list[str]] | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        cols = (columns or {}).get(name) or result_columns(rows)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="E11D48")
        for r in rows:
            vals = []
            for c in cols:
                v = r.get(c, "")
                if isinstance(v, float) and not math.isfinite(v):
                    v = None
                elif isinstance(v, (np.floating, np.integer)):
                    v = v.item()
                    if isinstance(v, float) and not math.isfinite(v):
                        v = None
                if isinstance(v, np.bool_):
                    v = bool(v)
                elif isinstance(v, str):
                    v = _ILLEGAL_XLSX.sub("", v)
                vals.append(v)
            ws.append(vals)
            for cell in ws[ws.max_row]:  # text that looks like a formula stays text (no formula injection)
                if isinstance(cell.value, str) and cell.value[:1] in "=+-@" and cell.value not in ("", "-"):
                    cell.data_type = "s"
        ws.freeze_panes = "B2"
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(str(c)) + 2))
    wb.save(path)


def write_table(rows: list[dict], path, columns: list[str] | None = None, sheet: str = "Results") -> Path:
    """Write rows × columns; the format follows the extension: .csv, .tsv / .txt (tab-separated), .xlsx, .slk (SYLK)
    or .dbf (dBase III)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xlsx":
        write_xlsx({sheet: rows}, path, {sheet: columns} if columns else None)
    elif ext == ".slk":
        write_sylk(rows, path, columns)
    elif ext == ".dbf":
        write_dbf(rows, path, columns)
    elif ext in (".tsv", ".txt", ".tab"):
        write_tsv(rows, path, columns)
    else:
        write_csv(rows, path, columns)
    return path


def table_text(rows: list[dict], columns: list[str], delimiter: str = "\t", header: bool = True,
               fmt=value_text) -> str:
    """Rows × columns as delimited text (for the clipboard); fmt(value) -> text (default: full precision)."""
    out = []
    if header:
        out.append(delimiter.join(str(c) for c in columns))
    for r in rows:
        out.append(delimiter.join(fmt(r.get(c)).replace(delimiter, " ").replace("\n", " ") for c in columns))
    return "\n".join(out) + "\n"


def results_workbook(project: Project, rows: list[dict], columns: list[str] | None = None,
                     segmented: bool = False) -> tuple[dict[str, list[dict]], dict[str, list[str]]]:
    """Sheets ({name: rows}) and their columns of a results workbook: whole-test results, time periods (segmented),
    zone visits of the tests in `rows`, animals, tests and settings (with the software version)."""
    whole = [r for r in rows if r.get("Period", "Whole test") == "Whole test"]
    sheets, cols = {"Results": whole}, {}
    main = [c for c in columns or [] if c != "Period"]
    if main:
        cols["Results"] = main
    seg = [r for r in rows if r.get("Period", "Whole test") != "Whole test"] if segmented else []
    if seg:
        sheets["Time periods"] = seg
        if main:
            cols["Time periods"] = columns if "Period" in columns else main[:1] + ["Period"] + main[1:]
    ids = {r.get("Test") for r in rows}
    visits = zone_visit_rows(project, [t for t in project.tests if t.id in ids])
    if visits:
        sheets["Zone visits"] = visits
    sheets["Animals"] = [{"Animal": a.id, "Group": a.group, "Sex": a.sex, **a.fields, "Notes": a.notes}
                         for a in project.animals]
    sheets["Tests"] = [{"Test": t.id, "Animal": t.animal_id, "Stage": t.stage, "Trial": t.trial, "Video": t.video,
                        "Apparatus": t.apparatus, "Start (s)": t.start_s, "Status": t.status, "User": t.experimenter,
                        "Reason for test end": t.end_reason, "Notes": t.notes}
                       for t in project.tests]
    sheets["Settings"] = [{"Setting": k, "Value": str(v)} for k, v in
                          {**{f"detection.{k}": v for k, v in project.detection.to_dict().items()},
                           **{f"analysis.{k}": v for k, v in project.analysis.to_dict().items()},
                           "test_duration_s": project.test_duration_s,
                           "software": f"mANY-MAZE {__version__}"}.items()]
    return sheets, cols


def export_results(project: Project, path, segmented: bool = False, columns: list[str] | None = None) -> Path:
    """Export project results; format chosen from the extension (.csv / .tsv / .txt / .xlsx / .slk / .dbf / .xml)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xml":
        return export_xml(project, path)
    rows = project.results(segmented=segmented)
    if ext == ".xlsx":
        sheets, cols = results_workbook(project, rows, columns, segmented)
        write_xlsx(sheets, path, cols)
    else:
        write_table(rows, path, columns)
    return path


def zone_visit_rows(project: Project, tests=None) -> list[dict]:
    """One row per zone visit (zone, entry time, exit time, duration) for every tracked test."""
    from .occupancy import zone_sequence

    rows = []
    for t in tests if tests is not None else project.tests:
        if t.status in INACTIVE_STATUSES or not project.has_track(t):
            continue
        app = project.apparatus_of(t)
        if app is None or not app.zones:
            continue
        s = project.analysis_for(t)
        tracks = project.load_tracks(t)
        ids = [t.animal_id] + list(t.extra_animals)
        for i, tr in enumerate(tracks):
            for zone, a, b in zone_sequence(tr, app, [z.name for z in app.zones], s):
                rows.append({"Test": t.id, "Animal": ids[i] if i < len(ids) else f"#{i + 1}", "Stage": t.stage,
                             "Zone": zone, "Entry (s)": round(a, 3), "Exit (s)": round(b, 3),
                             "Duration (s)": round(b - a, 3)})
    return rows


def export_animals(project: Project, path):
    """The animal list as CSV: ID, treatment, sex, the custom fields and the notes."""
    with atomic_write(path, newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["ID", "Treatment", "Sex"] + list(project.animal_fields) + ["Notes"])
        for a in project.animals:
            w.writerow([a.id, a.group, a.sex] + [a.fields.get(f, "") for f in project.animal_fields] + [a.notes])


def event_log_rows(project: Project, test) -> list[dict]:
    """Chronological log of a test (ANY-maze test event log): zone entries and exits (zones and zone groups),
    scored keys, I/O changes and pauses. Times are test times in seconds."""
    from .occupancy import zone_sequence

    rows = []
    ids = [test.animal_id] + list(test.extra_animals)
    app = project.apparatus_of(test)
    if app is not None and project.has_track(test):
        s = project.analysis_for(test)
        for i, tr in enumerate(project.load_tracks(test)):
            who = ids[i] if i < len(ids) else f"#{i + 1}"
            end = float(tr.t[-1] + tr.frame_durations()[-1]) if len(tr) else 0.0
            for zone, a, b in zone_sequence(tr, app, app.names(), s):
                rows.append({"Time (s)": a, "Animal": who, "Event": "Zone entry", "Detail": zone})
                if b < end - 1e-9:
                    rows.append({"Time (s)": b, "Animal": who, "Event": "Zone exit", "Detail": zone})
    kinds = {b.name: b.kind for b in project.behaviours}
    for e in test.events or []:
        name = e.get("behaviour", "")
        if e.get("t_end") is None or kinds.get(name) == "point":
            rows.append({"Time (s)": float(e["t"]), "Animal": test.animal_id, "Event": "Key", "Detail": name})
        else:
            rows.append({"Time (s)": float(e["t"]), "Animal": test.animal_id, "Event": "Key on", "Detail": name})
            rows.append({"Time (s)": float(e["t_end"]), "Animal": test.animal_id, "Event": "Key off",
                         "Detail": name})
    for e in test.io_events or []:
        kind = {"input": "Input", "variable": "Variable"}.get(e.get("kind"), "Output")
        ch = ":".join(str(x) for x in (e.get("device"), e.get("channel")) if x not in (None, ""))
        rows.append({"Time (s)": float(e.get("t", 0.0)), "Animal": test.animal_id, "Event": kind,
                     "Detail": f"{ch} = {value_text(e.get('value'))}"})
    for pz in test.pauses or []:
        rows.append({"Time (s)": float(pz[0]), "Animal": test.animal_id, "Event": "Paused", "Detail": ""})
        if len(pz) > 1 and pz[1] is not None and float(pz[1]) > float(pz[0]):
            rows.append({"Time (s)": float(pz[1]), "Animal": test.animal_id, "Event": "Resumed", "Detail": ""})
    order = {"Zone exit": 0, "Key off": 1, "Resumed": 2}
    rows.sort(key=lambda r: (r["Time (s)"], order.get(r["Event"], 5)))
    for r in rows:
        r["Time (s)"] = round(r["Time (s)"], 3)
    return rows


def wide_rows(rows: list[dict], measures: list[str], across: tuple[str, ...] = ("Stage", "Trial"),
              keep: tuple[str, ...] = ("Animal", "Group", "Sex")) -> list[dict]:
    """One row per animal with one column per measure × stage/trial (ANY-maze "one row per animal" layout,
    ready for Prism or repeated-measures analysis elsewhere). Rows of a time-period analysis are kept apart by
    their period. Columns are named "<measure> [Day 1 · 2]"."""
    def label(r):
        parts = [str(r.get(k, "")) for k in across if r.get(k, "") not in ("", None)]
        per = r.get("Period", "")
        if per and per != "Whole test":
            parts.append(str(per))
        return " · ".join(parts)

    out: dict = {}
    levels: list[str] = []
    for r in rows:
        key = r.get("Animal", "")
        row = out.setdefault(key, {k: r.get(k, "") for k in keep if k in r})
        lab = label(r)
        if lab not in levels:
            levels.append(lab)
        for m in measures:
            col = f"{m} [{lab}]" if lab else m
            if m in r and col not in row:  # first attempt wins (superseded attempts are already excluded)
                row[col] = r[m]
    return list(out.values())


def trial_means(rows: list[dict], measures: list[str], keep: tuple[str, ...] = ("Group", "Sex")) -> list[dict]:
    """Mean of each animal's trials per stage (and time period): one row per animal × stage, with the number of
    trials averaged ("Trials"). Text measures are left out. Water-maze style blocks of trials per day."""
    groups: dict = {}
    for r in rows:
        key = (r.get("Animal", ""), r.get("Stage", ""), r.get("Period", "Whole test"))
        groups.setdefault(key, []).append(r)
    out = []
    for (animal, stage, period), rs in groups.items():
        row = {"Animal": animal, **{k: rs[0].get(k, "") for k in keep if k in rs[0]}, "Stage": stage}
        if period and period != "Whole test":
            row["Period"] = period
        row["Trials"] = len({r.get("Test") for r in rs})
        for m in measures:
            vals = [float(r[m]) for r in rs if isinstance(r.get(m), (int, float, np.integer, np.floating))
                    and not isinstance(r.get(m), bool) and math.isfinite(float(r[m]))]
            if vals:
                row[m] = round(float(np.mean(vals)), 4)
        out.append(row)
    return out


def protocol_report(project: Project, path) -> Path:
    """Self-contained HTML description of the protocol (ANY-maze protocol report): experiment, stages, keys,
    apparatus maps with every zone / point / line / group / sequence, animal tracking and analysis settings,
    procedures, I/O devices and training criteria."""
    from . import plots
    from .measures import AnalysisSettings
    from .procedures import describe_statement, normalize_procedures
    from .tracking import DetectionSettings
    from .workflow import criterion_text

    def table(header, body):
        h = "".join(f"<th>{html.escape(str(c))}</th>" for c in header)
        b = "".join("<tr>" + "".join(f"<td>{html.escape(value_text(v))}</td>" for v in r) + "</tr>" for r in body)
        return f"<table><tr>{h}</tr>{b}</table>"

    def settings(obj, default):
        d, d0 = obj.to_dict(), default.to_dict()
        return table(("Setting", "Value", ""), [(k.replace("_", " ").capitalize(), v if v != [] else "—",
                                                 "" if v == d0.get(k) else "changed") for k, v in d.items()])

    css = """
    body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Helvetica,Arial,sans-serif;margin:24px;color:#0f172a}
    h1{color:#e11d48} table{border-collapse:collapse;font-size:12px;margin-bottom:10px}
    td,th{border:1px solid #e2e8f0;padding:3px 6px;text-align:left} th{background:#f1f5f9}
    .card{border:1px solid #e2e8f0;border-radius:8px;padding:8px;margin-bottom:12px} pre{background:#f8fafc;padding:8px}
    """
    p = project
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(p.name)} – protocol</title>"
           f"<style>{css}</style></head><body><h1>{html.escape(p.name)}: protocol</h1>"]
    if p.description:
        out.append(f"<p>{html.escape(p.description)}</p>")
    out.append(f"<p>Generated {_dt.datetime.now():%Y-%m-%d %H:%M} by mANY-MAZE {__version__}.</p>")
    out.append("<h2>Protocol</h2>" + table(("Item", "Value"), [
        ("Protocol", p.protocol), ("Test duration (s)", p.test_duration_s or "until the end of the video"),
        ("Test starts", {"on_detection": "when the animal is first detected",
                         "experimenter_leaves": "when the experimenter's hand has left the image"}.get(
            p.start_mode, "at the test's start time")),
        ("Blind testing", "yes" if p.blind else "no"),
        ("Confirm the animal's ID", "yes" if p.settings_extra.get("confirm_id") else "no"),
        ("Stages", ", ".join(p.stages) or "—"), ("Treatments", ", ".join(g.name for g in p.groups) or "—"),
        ("Animal columns", ", ".join(p.animal_fields) or "—"),
        ("Animals / tests", f"{len(p.animals)} / {len(p.tests)}")]))
    if p.behaviours:
        out.append("<h2>Keys</h2>" + table(("Behaviour", "Key", "Type", "Exclusive set"),
                                           [(b.name, b.key, b.kind, b.group) for b in p.behaviours]))
    for app in p.apparatus:
        out.append(f"<h2>Apparatus: {html.escape(app.name)}</h2><div class='card'>")
        try:
            from matplotlib.figure import Figure

            fig = Figure(figsize=(4.2, 4.2))
            ax = fig.add_subplot(111)
            plots._draw_apparatus(ax, app, labels=True)
            plots._limits(ax, app, None)
            ax.set_xticks([])
            ax.set_yticks([])
            out.append(_img(plots.fig_to_png(fig), 380))
        except Exception as e:
            log.warning("apparatus figure for %s failed: %s", app.name, e)
            out.append("<p><em>(figure unavailable)</em></p>")
        cal = (f"{app.px_per_cm:.3f} px/cm" + (f" (line of {app.calibration_length_cm:g} cm)"
                                                 if app.calibration_length_cm else "")) if app.px_per_cm else \
            "not calibrated (results in pixels)"
        out.append(table(("Item", "Value"), [("Template", app.template), ("Calibration", cal),
                                             ("Arena", app.arena.to_dict()["type"] if app.arena else "—")]))
        if app.zones:
            out.append(table(("Zone", "Shape", f"Area ({app.unit}²)", "Entry rule", "Options"), [
                (z.name, z.shape.to_dict()["type"], round(z.shape.area() * app.scale ** 2, 2),
                 ENTRY_RULE_TEXT.get(z.entry_rule, z.entry_rule),
                 ", ".join(x for x, on in (("hidden", z.hidden), ("moveable", z.moveable),
                                           (f"investigate {z.investigation_distance_cm:g} cm",
                                            z.investigation_distance_cm > 0),
                                           (f"entry facing the zone (±{z.entry_orientation_deg:g}°)",
                                            z.entry_orientation_deg > 0),
                                           (f"Whishaw's corridor {z.whishaw_width_cm:g} {app.unit}",
                                            z.whishaw_width_cm > 0)) if on))
                for z in app.zones]))
        if app.groups:
            out.append(table(("Zone group", "Zones", "Excluding"),
                             [(g.name, ", ".join(g.zones), ", ".join(g.exclude)) for g in app.groups]))
        if app.points:
            out.append(table(("Point", "x (px)", "y (px)", "Radius (cm)"),
                             [(q.name, round(q.x, 1), round(q.y, 1), q.radius_cm) for q in app.points]))
        if app.lines:
            out.append(table(("Line", "From (px)", "To (px)"), [
                (ln.name, f"{ln.x1:.0f}, {ln.y1:.0f}", f"{ln.x2:.0f}, {ln.y2:.0f}") for ln in app.lines]))
        if app.sequences:
            out.append(table(("Sequence", "Steps", "Options"), [
                (q.name, " → ".join(q.steps), ", ".join(x for x, on in (
                    ("must begin at the first step", q.from_start), ("other zones allowed", q.allow_other),
                    ("both directions", q.bidirectional), ("overlapping", q.overlap),
                    (f"time limit {q.max_duration_s:g} s", q.max_duration_s > 0)) if on))
                for q in app.sequences]))
        out.append("</div>")
    out.append("<h2>Animal tracking</h2>" + settings(p.detection, DetectionSettings()))
    out.append("<h2>Analysis</h2>" + settings(p.analysis, AnalysisSettings()))
    procs = normalize_procedures(p.procedures)
    if procs:
        out.append("<h2>Procedures</h2>")
        for pr in procs:

            def lines(stmts, depth=0):
                res = []
                for st in stmts or []:
                    res.append("  " * depth + ("" if st.get("enabled", True) else "[off] ") + describe_statement(st))
                    for k in ("body", "else"):
                        if isinstance(st.get(k), list):
                            if k == "else" and st[k]:
                                res.append("  " * depth + "else")
                            res += lines(st[k], depth + 1)
                return res

            state = "" if pr.get("enabled", True) else " (disabled)"
            out.append(f"<div class='card'><b>{html.escape(pr.get('name', 'Procedure'))}{state}</b>"
                       f"<pre>{html.escape(chr(10).join(lines(pr.get('statements'))))}</pre></div>")
    if p.io_devices:
        out.append("<h2>I/O devices</h2>" + table(("Device", "Type", "Settings"), [
            (d.get("name", ""), d.get("type", ""), ", ".join(f"{k}={v}" for k, v in d.items()
                                                             if k not in ("name", "type") and not isinstance(v, (list, dict))))
            for d in p.io_devices]))
    if p.training_criteria:
        out.append("<h2>Training criteria</h2><ul>" + "".join(
            f"<li>{html.escape(criterion_text(c))}</li>" for c in p.training_criteria) + "</ul>")
    out.append("</body></html>")
    path = Path(path)
    write_text_atomic(path, "\n".join(out))
    return path


def export_track(track, path):
    track.to_csv(path)


# ---------------------------------------------------------------- per-test raw data
def export_raw_data(project: Project, out_dir, tests=None, parameters: list[str] | None = None,
                    derived: bool = True, delimiter: str = ",", progress=None, should_stop=None) -> list[Path]:
    """One CSV per test and animal: time, raw track columns and (derived=True) every per-frame parameter of
    core.charts (or only `parameters`). Returns the written paths."""
    from . import charts
    from .track import COLUMNS

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    tests = [t for t in (tests if tests is not None else project.tests) if project.has_track(t)]
    written = []
    beh = project.behaviours
    for k, t in enumerate(tests):
        if should_stop and should_stop():
            break
        tracks = project.load_tracks(t)
        app = project.apparatus_of(t)
        ids = [t.animal_id] + list(t.extra_animals)
        for i, tr in enumerate(tracks):
            raw_cols = [c for c in COLUMNS if c != "t"]
            cols = ["Time (s)"] + [f"raw {c}" for c in raw_cols]
            arrays = [tr.t] + [getattr(tr, c).astype(float) for c in raw_cols]
            if derived and app is not None and len(tr):
                others = [o for j, o in enumerate(tracks) if j != i]
                names = parameters
                if names is not None:
                    avail = {p.name for p in charts.parameters(app, tr, beh, len(others))}
                    names = [n for n in names if n in avail]
                dcols, arr = charts.per_frame_table(tr, app, project.analysis_for(t), t.events if i == 0 else [],
                                                    beh, names, others or None)
                cols += dcols[1:]
                arrays += [arr[:, j] for j in range(1, arr.shape[1])]
            aid = ids[i] if i < len(ids) else f"{t.animal_id}#{i + 1}"
            safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in str(aid))
            p = out_dir / f"test_{t.id:04d}_{safe}.{'tsv' if delimiter == chr(9) else 'csv'}"
            with atomic_write(p, newline="") as f:
                f.write(f"# Test {t.id}, animal {aid}, stage {t.stage}, trial {t.trial}, "
                        f"unit {app.unit if app else 'px'} (raw columns in pixels)\n")
                w = csv.writer(f, delimiter=delimiter)
                w.writerow(cols)
                M = np.column_stack(arrays) if arrays and len(tr) else np.zeros((0, len(cols)))
                for row in M:
                    # Time with fixed decimals: .6g would round it to 0.1 s past 10000 s
                    w.writerow(["" if not np.isfinite(v) else (f"{v:.5f}" if j == 0 else f"{v:.6g}")
                                for j, v in enumerate(row)])
            written.append(p)
        if progress:
            progress((k + 1) / max(1, len(tests)))
    return written


# ---------------------------------------------------------------- XML export
def _attrs(**kw) -> str:
    parts = []
    for k, v in kw.items():
        if v is None:
            continue
        if isinstance(v, bool):
            v = "true" if v else "false"
        elif isinstance(v, (float, np.floating)):
            v = "" if not math.isfinite(v) else f"{float(v):.10g}"
        parts.append(f"{k.replace('_', '-')}={_qa(str(v))}")
    return (" " + " ".join(parts)) if parts else ""


def _shape_xml(tag: str, d: dict | None, ind: str) -> str:
    if not d:
        return ""
    scal = {k: v for k, v in d.items() if k not in ("points", "type")}
    pts = d.get("points") or []
    if not pts:
        return f"{ind}<{tag}{_attrs(type=d.get('type'), **scal)}/>\n"
    inner = "".join(f"{ind}  <vertex{_attrs(x=float(p[0]), y=float(p[1]))}/>\n" for p in pts)
    return f"{ind}<{tag}{_attrs(type=d.get('type'), **scal)}>\n{inner}{ind}</{tag}>\n"


def _value_attrs(v) -> dict:
    if isinstance(v, (bool, np.bool_)):
        return {"value": bool(v), "type": "bool"}
    if isinstance(v, (int, np.integer)):
        return {"value": int(v), "type": "integer"}
    if isinstance(v, (float, np.floating)):
        return {"value": float(v), "type": "number"}
    if isinstance(v, (list, tuple, dict)):
        import json

        return {"value": json.dumps(v, default=str), "type": "json"}
    return {"value": "" if v is None else str(v), "type": "text"}


def export_xml(project: Project, path, tests=None, include_tracks: bool = True, include_results: bool = True,
               rows: list[dict] | None = None, segmented: bool = False, progress=None, should_stop=None) -> Path:
    """Write the whole experiment (settings, apparatus, animals, tests with scoring, I/O events, results and the
    per-frame tracking data) as one XML document. Written incrementally, one test at a time.

    rows: precomputed result rows (e.g. from the Results page); otherwise results are calculated per test.
    Track columns are space-separated numbers ("NaN" = missing) so they load with str2num / sscanf in MATLAB.
    """
    from .track import COLUMNS

    path = Path(path)
    tests = [t for t in (tests if tests is not None else project.tests)]
    by_test: dict = {}
    if rows is not None:
        for r in rows:
            by_test.setdefault(r.get("Test"), []).append(r)
    tmp = path.with_name(path.name + ".part")
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            w = f.write
            w('<?xml version="1.0" encoding="UTF-8"?>\n')
            w(f"<manymaze-experiment{_attrs(format_version=XML_FORMAT_VERSION, software=f'mANY-MAZE {__version__}', exported=_dt.datetime.now().isoformat(timespec='seconds'))}>\n")
            w(f"  <experiment{_attrs(name=project.name, protocol=project.protocol, test_duration_s=float(project.test_duration_s), start_mode=project.start_mode, created=project.created, blind=bool(getattr(project, 'blind', False)))}>\n")
            w(f"    <description>{_xesc(project.description or '')}</description>\n")
            for tag, d in (("detection-settings", project.detection.to_dict()),
                           ("analysis-settings", project.analysis.to_dict())):
                w(f"    <{tag}>\n")
                for k, v in d.items():
                    w(f"      <setting{_attrs(name=k, **_value_attrs(v))}/>\n")
                w(f"    </{tag}>\n")
            variables = project.variables
            if variables:
                w("    <variables>\n")
                for k, v in variables.items():
                    w(f"      <variable{_attrs(name=k, **_value_attrs(v))}/>\n")
                w("    </variables>\n")
            w("  </experiment>\n")
            w("  <groups>\n" + "".join(f"    <group{_attrs(name=g.name, color=g.color)}/>\n" for g in project.groups)
              + "  </groups>\n")
            w("  <stages>\n" + "".join(f"    <stage{_attrs(name=s)}/>\n" for s in project.stages) + "  </stages>\n")
            if project.experimenters:
                w("  <experimenters>\n" + "".join(f"    <experimenter{_attrs(name=u)}/>\n" for u in project.experimenters)
                  + "  </experimenters>\n")
            w("  <behaviours>\n" + "".join(f"    <behaviour{_attrs(name=b.name, key=b.key, kind=b.kind)}/>\n"
                                           for b in project.behaviours) + "  </behaviours>\n")
            w("  <apparatus-list>\n")
            for a in project.apparatus:
                fs = a.frame_size or (None, None)
                w(f"    <apparatus{_attrs(name=a.name, template=a.template, unit=a.unit, px_per_cm=a.px_per_cm, frame_width=fs[0], frame_height=fs[1])}>\n")
                w(_shape_xml("arena", a.arena.to_dict() if a.arena else None, "      "))
                for z in a.zones:
                    w(f"      <zone{_attrs(name=z.name, color=z.color)}>\n")
                    w(_shape_xml("shape", z.shape.to_dict(), "        "))
                    w("      </zone>\n")
                for g in a.groups:
                    w(f"      <zone-group{_attrs(name=g.name)}>\n")
                    for zn in g.zones:
                        w(f"        <member{_attrs(zone=zn)}/>\n")
                    for zn in g.exclude:
                        w(f"        <exclude{_attrs(zone=zn)}/>\n")
                    w("      </zone-group>\n")
                for p in a.points:
                    w(f"      <point{_attrs(name=p.name, x=p.x, y=p.y, radius_cm=p.radius_cm, color=p.color)}/>\n")
                for ln in a.lines:
                    w(f"      <line{_attrs(name=ln.name, x1=ln.x1, y1=ln.y1, x2=ln.x2, y2=ln.y2, color=ln.color)}/>\n")
                w("    </apparatus>\n")
            w("  </apparatus-list>\n")
            w("  <animals>\n")
            for an in project.animals:
                w(f"    <animal{_attrs(id=an.id, group=an.group, sex=an.sex, notes=an.notes or None)}")
                if an.fields:
                    w(">\n" + "".join(f"      <field{_attrs(name=k, **_value_attrs(v))}/>\n" for k, v in an.fields.items())
                      + "    </animal>\n")
                else:
                    w("/>\n")
            w("  </animals>\n")
            w(f"  <tests{_attrs(count=len(tests))}>\n")
            for k, t in enumerate(tests):
                if should_stop and should_stop():
                    break
                w(f"    <test{_attrs(id=t.id, animal=t.animal_id, stage=t.stage, trial=t.trial, apparatus=t.apparatus, video=t.video, start_s=float(t.start_s), duration_s=float(t.duration_s or project.test_duration_s), status=t.status, recorded_at=t.recorded_at, experimenter=t.experimenter or None, end_reason=t.end_reason or None)}>\n")
                for ea in t.extra_animals:
                    w(f"      <extra-animal{_attrs(id=ea)}/>\n")
                if t.notes:
                    w(f"      <notes>{_xesc(t.notes)}</notes>\n")
                if t.variables:
                    w("      <variables>\n" + "".join(f"        <variable{_attrs(name=n, **_value_attrs(v))}/>\n"
                                                     for n, v in t.variables.items()) + "      </variables>\n")
                ov = dict(t.zone_overrides)
                pos = ov.pop(POSITION_KEY, None)
                if isinstance(pos, dict):
                    w(f"      <apparatus-position{_attrs(**position_args(pos))}/>\n")
                cal = ov.pop(CALIBRATION_KEY, None)
                if isinstance(cal, dict):
                    x1, y1, x2, y2 = cal.get("calibration_line") or [None] * 4
                    w(f"      <calibration{_attrs(px_per_cm=cal.get('px_per_cm'), length_cm=cal.get('calibration_length_cm'), x1=x1, y1=y1, x2=x2, y2=y2)}/>\n")
                if ov:
                    w("      <zone-overrides>\n")
                    for zn, sd in ov.items():
                        w(f"        <zone{_attrs(name=zn)}>\n" + _shape_xml("shape", sd, "          ") + "        </zone>\n")
                    w("      </zone-overrides>\n")
                pauses = t.pauses
                if pauses:
                    # a pause still open (no end: [t, None] or [t]) is written with an empty end
                    w("      <pauses>\n" + "".join(
                        f"        <pause{_attrs(start=float(pz[0]), end=float(pz[1]) if len(pz) > 1 and pz[1] is not None else '')}/>\n"
                        for pz in pauses) + "      </pauses>\n")
                w("      <events>\n" + "".join(
                    f"        <event{_attrs(behaviour=e.get('behaviour'), t=float(e['t']), t_end=None if e.get('t_end') is None else float(e['t_end']))}/>\n"
                    for e in t.events) + "      </events>\n")
                io = t.io_events
                if io:
                    w("      <io-events>\n" + "".join(
                        f"        <io{_attrs(t=float(e.get('t', 0)), device=e.get('device'), channel=e.get('channel'), kind=e.get('kind'), value=e.get('value'))}/>\n"
                        for e in io) + "      </io-events>\n")
                rv = t.result_variables
                if rv:
                    w("      <result-variables>\n" + "".join(f"        <variable{_attrs(name=n, **_value_attrs(v))}/>\n"
                                                            for n, v in rv.items()) + "      </result-variables>\n")
                has = project.has_track(t)
                if include_results and t.status not in INACTIVE_STATUSES:
                    trows = by_test.get(t.id) if rows is not None else None
                    if trows is None and project.has_results(t):
                        try:
                            trows = project.analyse_test(t, segmented)
                        except Exception as e:
                            log.warning("analysis of test %s failed in the XML export: %s", t.id, e)
                            trows = []
                            w(f"      <results-error{_attrs(message=str(e))}/>\n")
                    info = {*INFO_COLUMNS, *project.animal_fields}
                    for r in trows or []:
                        w(f"      <results{_attrs(animal=r.get('Animal'), period=r.get('Period', 'Whole test'))}>\n")
                        for c, v in r.items():
                            if c in info:
                                continue
                            w(f"        <result{_attrs(name=c, **_value_attrs(v))}/>\n")
                        w("      </results>\n")
                if include_tracks and has:
                    ids = [t.animal_id] + list(t.extra_animals)
                    for i, tr in enumerate(project.load_tracks(t)):
                        aid = ids[i] if i < len(ids) else f"{t.animal_id}#{i + 1}"
                        w(f"      <track{_attrs(animal=aid, index=i + 1, fps=float(tr.fps), samples=len(tr), video_start_s=tr.meta.get('video_start_s'), units='px')}>\n")
                        for c in COLUMNS:
                            v = getattr(tr, c).astype(float)
                            txt = " ".join("NaN" if not np.isfinite(x) else (f"{x:.5f}" if c == "t" else f"{x:.3f}")
                                           for x in v) if c != "detected" else " ".join("1" if x else "0" for x in v)
                            w(f"        <column{_attrs(name=c)}>{txt}</column>\n")
                        w("      </track>\n")
                w("    </test>\n")
                if progress:
                    progress((k + 1) / max(1, len(tests)))
            w("  </tests>\n</manymaze-experiment>\n")
    except BaseException:  # no half-written .part left behind
        tmp.unlink(missing_ok=True)
        raise
    if should_stop and should_stop():
        tmp.unlink(missing_ok=True)
        return None
    tmp.replace(path)
    return path


def read_experiment_xml(path) -> dict:
    """Parse an XML export back into plain Python data (tests with results, events and track arrays)."""
    import xml.etree.ElementTree as ET

    def val(e):
        v, t = e.get("value", ""), e.get("type", "text")
        if t == "number":
            return float(v) if v != "" else math.nan
        if t == "integer":
            return int(v)
        if t == "bool":
            return v == "true"
        if t == "json":
            import json

            return json.loads(v)
        return v

    root = ET.parse(path).getroot()
    exp = root.find("experiment")
    out = {"name": exp.get("name"), "format_version": int(root.get("format-version", "1")),
           "animals": [dict(a.attrib) for a in root.iter("animal")],
           "experimenters": [e.get("name") for e in root.iter("experimenter")],
           "apparatus": [a.get("name") for a in root.find("apparatus-list")], "tests": []}
    for t in root.find("tests"):
        d = dict(t.attrib)
        d["events"] = [dict(e.attrib) for e in t.findall("events/event")]
        d["io_events"] = [dict(e.attrib) for e in t.findall("io-events/io")]
        d["results"] = [{"animal": r.get("animal"), "period": r.get("period"),
                         "values": {x.get("name"): val(x) for x in r.findall("result")}} for r in t.findall("results")]
        d["tracks"] = []
        for tr in t.findall("track"):
            cols = {c.get("name"): np.array(c.text.split(), float) if c.text else np.zeros(0)
                    for c in tr.findall("column")}
            d["tracks"].append({"animal": tr.get("animal"), "fps": float(tr.get("fps")), "columns": cols})
        out["tests"].append(d)
    return out


# ---------------------------------------------------------------- HTML report
def _img(png: bytes, width=320) -> str:
    return f'<img width="{width}" src="data:image/png;base64,{base64.b64encode(png).decode()}">'


def html_report(project: Project, path, tests=None, include_plots: bool = True, measures: list[str] | None = None,
                stats_measures: list[str] | None = None, heatmap_norm: str = "auto",
                chart_parameters: list[str] | None = None, color_by: str = "time", rows: list[dict] | None = None
                ) -> Path:
    """Self-contained HTML report: summary, per-test track plots/heat maps (+ optional charts of per-frame
    parameters), group heat maps on a common scale, results and statistics (compared between treatments as the
    Statistics page does). rows: the results to tabulate and compare (default: the whole-test results of `tests`)."""
    from . import analyses, charts, plots

    tests = tests if tests is not None else [t for t in project.tests
                                             if t.status not in INACTIVE_STATUSES and project.has_results(t)]
    if rows is None:
        rows = [r for t in tests for r in project.analyse_test(t)]
    cols = result_columns(rows)
    if measures:
        info = ["Test", "Animal", "Group", "Stage", "Trial"]
        cols = [c for c in info if c in cols] + [c for c in measures if c in cols]
    css = """
    body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Helvetica,Arial,sans-serif;margin:24px;color:#0f172a}
    h1{color:#e11d48} table{border-collapse:collapse;font-size:12px} td,th{border:1px solid #e2e8f0;padding:3px 6px}
    th{background:#f1f5f9;position:sticky;top:0} .grid{display:flex;flex-wrap:wrap;gap:12px}
    .card{border:1px solid #e2e8f0;border-radius:8px;padding:8px} pre{background:#f8fafc;padding:8px}
    """
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(project.name)}</title>"
           f"<style>{css}</style></head><body>"]
    out.append(f"<h1>{html.escape(project.name)}</h1>")
    out.append(f"<p>{html.escape(project.description)}</p>")
    out.append(f"<p>Generated {_dt.datetime.now():%Y-%m-%d %H:%M} by mANY-MAZE {__version__}. "
               f"{len(tests)} tests, {len(project.animals)} animals.</p>")
    beh = project.behaviours
    if include_plots and tests:
        out.append("<h2>Tracks</h2><div class='grid'>")
        hm_vmax = None
        if heatmap_norm == "fixed":
            hm_vmax = 0.0
            for t in tests:
                trs = project.load_tracks(t)
                if trs:
                    H, _ = plots.occupancy(trs[:1], project.apparatus_of(t))
                    hm_vmax = max(hm_vmax, float(H.max()) if H.size else 0.0)
            hm_vmax = hm_vmax or None
        for t in tests:
            tracks = project.load_tracks(t)
            if not tracks:
                continue
            app = project.apparatus_of(t)
            frame = project.start_frame(t)
            an = project.get_animal(t.animal_id)
            title = f"Test {t.id} · {t.animal_id} · {an.group if an else ''}"
            markers = plots.behaviour_markers(tracks[0], app, project.analysis_for(t), t.events, beh)
            tp = plots.fig_to_png(plots.track_plot(tracks[0], app, frame=frame, size=(3.2, 3.2), color_by=color_by,
                                                   markers=markers, settings=project.analysis_for(t)))
            hm = plots.fig_to_png(plots.heatmap(tracks[:1], app, size=(3.6, 3.2), vmax=hm_vmax,
                                                norm="auto" if heatmap_norm == "fixed" else heatmap_norm))
            card = f"<div class='card'><b>{html.escape(title)}</b><br>{_img(tp, 260)}{_img(hm, 290)}"
            if chart_parameters:
                names = [p.name for p in charts.parameters(app, tracks[0], beh) if p.name in chart_parameters]
                if names:
                    fig = plots.chart_figure(tracks[0], app, names, project.analysis_for(t), t.events, beh,
                                              size=(6.5, 1.0 + 1.1 * len(names)))
                    card += "<br>" + _img(plots.fig_to_png(fig), 560)
            out.append(card + "</div>")
        out.append("</div>")
        # group heatmaps
        groups = {}
        for t in tests:
            a = project.get_animal(t.animal_id)
            groups.setdefault(a.group if a and a.group else "No group", []).append(t)
        if len(groups) > 1:
            fig = plots.group_heatmap(project, groups, norm="auto" if heatmap_norm == "fixed" else heatmap_norm)
            out.append("<h2>Group occupancy</h2><div class='grid'><div class='card'>"
                       f"{_img(plots.fig_to_png(fig), 300 * min(3, len(groups)) + 60)}</div></div>")
    if stats_measures and rows:
        out.append("<h2>Statistics</h2>")
        for m in stats_measures:
            a = analyses.compare(project, rows, m, "Group")
            if len(a.result["groups"]) < 2:
                continue
            out.append(f"<div class='card'>{_img(plots.fig_to_png(a.figure), 360)}"
                       f"<pre>{html.escape(a.summary_text)}</pre></div>")
    out.append("<h2>Results</h2><table><tr>" + "".join(f"<th>{html.escape(str(c))}</th>" for c in cols) + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{html.escape(value_text(r.get(c)))}</td>" for c in cols) + "</tr>")
    out.append("</table></body></html>")
    path = Path(path)
    write_text_atomic(path, "\n".join(out))
    return path
