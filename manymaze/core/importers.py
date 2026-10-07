"""Import data exported by ANY-maze (or other tracking software) as spreadsheets.

ANY-maze can save its spreadsheets (animals, test schedule, results) and per-test track data as CSV, tab-separated
text or Excel files. Their exact columns depend on what the user selected, so imports work by mapping columns to
roles; :func:`guess_mapping` proposes a mapping from the header names (ANY-maze, EthoVision and generic names).
"""

from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path

import numpy as np

from .track import Track

# role -> header patterns (lower case, regular expressions), most specific first
ANIMAL_ROLES = {
    "id": [r"^animal id$", r"^animal$", r"^subject( id)?$", r"^id$", r"^animal (no|number|#)$",
           r"^(rat|mouse)( id| no| number)?$"],
    "group": [r"^treatment", r"^group", r"^condition", r"^genotype"],
    "sex": [r"^sex$", r"^gender$"],
}
TEST_ROLES = {
    "test": [r"^test( no| number| #)?$", r"^trial id$"],
    "animal": [r"^animal id$", r"^animal$", r"^subject"],
    "group": [r"^treatment", r"^group"],
    "stage": [r"^stage", r"^day$", r"^session$"],
    "trial": [r"^trial$", r"^trial (no|number|#)"],
    "apparatus": [r"^apparatus", r"^arena$", r"^maze$"],
    "video": [r"video", r"file ?name", r"^source"],
    "duration": [r"^duration", r"^test duration"],
}
TRACK_ROLES = {
    "t": [r"^time", r"^t$", r"^recording time", r"^trial time"],
    "x": [r"^cent(re|er)[^,]*\bx\b", r"cent(re|er).*position.*x", r"^x$", r"^x cent", r"^x pos", r"^position x",
          r"^body.*x"],
    "y": [r"^cent(re|er)[^,]*\by\b", r"cent(re|er).*position.*y", r"^y$", r"^y cent", r"^y pos", r"^position y",
          r"^body.*y"],
    "hx": [r"^head.*\bx\b", r"^nose.*\bx\b", r"^x nose", r"^x head"],
    "hy": [r"^head.*\by\b", r"^nose.*\by\b", r"^y nose", r"^y head"],
    "tx": [r"^tail.*\bx\b", r"^x tail"],
    "ty": [r"^tail.*\by\b", r"^y tail"],
}


def _read_text(p: Path) -> str:
    """Decode a text export: UTF-16/32 by BOM (Excel "Unicode text"), else UTF-8, else Windows cp1252."""
    raw = p.read_bytes()
    for bom, enc in ((b"\xff\xfe\x00\x00", "utf-32"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if raw.startswith(bom):
            return raw.decode(enc)
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def read_table(path: str | Path) -> tuple[list[str], list[list[str]]]:
    """Header and rows of a CSV / TSV / text / Excel file (first sheet). Values are strings."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(p, read_only=True, data_only=True).worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = _read_text(p)
        sample = text[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in sample else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text, newline=""), dialect)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    header = [h.strip() for h in rows[0]]
    n = len(header)
    body = [(r + [""] * n)[:n] for r in rows[1:]]
    return header, body


def guess_mapping(header: list[str], roles: dict[str, list[str]]) -> dict[str, int | None]:
    """Column index for each role (None = not found), matching header names against the role patterns."""
    low = [re.sub(r"\s+", " ", h.lower().replace("_", " ")).strip() for h in header]
    used: set[int] = set()
    out: dict[str, int | None] = {}
    for role, pats in roles.items():
        out[role] = None
        for pat in pats:
            hit = next((i for i, h in enumerate(low) if i not in used and re.search(pat, h)), None)
            if hit is not None:
                out[role] = hit
                used.add(hit)
                break
    return out


_TIME_UNITS = {"ms": 0.001, "min": 60.0, "mins": 60.0, "h": 3600.0, "hr": 3600.0, "hrs": 3600.0}


def parse_number(s: str) -> float:
    """Numbers as ANY-maze writes them: "2.168m", "18.9s", "0:43", "1,5", "" -> nan."""
    s = str(s).strip()
    if not s:
        return math.nan
    if re.fullmatch(r"[-+]?\d+:\d{1,2}(:\d{1,2})?(\.\d+)?", s):  # h:mm:ss / m:ss
        parts = [float(x) for x in s.lstrip("-+").split(":")]
        v = 0.0
        for x in parts:
            v = v * 60 + x
        return -v if s.startswith("-") else v
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):  # a date/time cell is not a number
        return math.nan
    compact = s.replace(" ", "")
    if re.fullmatch(r"[-+]?\d{1,3}(,\d{3})+(\.\d+)?", compact):  # 1,234.5 — thousands separators
        compact = compact.replace(",", "")
    m = re.match(r"^[-+]?(\d+[.,]?\d*|[.,]\d+)([eE][-+]?\d+)?", compact)
    if not m:
        return math.nan
    v = float(m.group(0).replace(",", "."))
    unit = compact[m.end():].lower()
    return v * _TIME_UNITS.get(unit, 1.0)  # "5 min" → 300 s, "150ms" → 0.15 s; other units are kept as-is


def _cell(row, idx):
    return row[idx].strip() if idx is not None and idx < len(row) else ""


def import_animals(project, header, rows, mapping: dict, extra_fields: list[int] | None = None) -> list[str]:
    """Create / update animals from a table. Other mapped columns become custom animal fields."""
    ids = []
    for r in rows:
        aid = _cell(r, mapping.get("id"))
        if not aid:
            continue
        group = _cell(r, mapping.get("group"))
        a = project.ensure_animal(aid, group)
        if group:
            a.group = group
        if mapping.get("sex") is not None:
            a.sex = _cell(r, mapping["sex"])
        for i in extra_fields or []:
            name = header[i]
            if name not in project.animal_fields:
                project.animal_fields.append(name)
            a.fields[name] = _cell(r, i)
        ids.append(aid)
    return ids


def import_tests(project, header, rows, mapping: dict, video_dir: str | Path | None = None) -> list:
    """Create tests (test schedule) from a table: animal, stage, trial, apparatus, video, duration."""
    new = []
    for r in rows:
        aid = _cell(r, mapping.get("animal"))
        if not aid:
            continue
        a = project.ensure_animal(aid, _cell(r, mapping.get("group")))
        if not a.group and _cell(r, mapping.get("group")):
            a.group = _cell(r, mapping.get("group"))
        stage = project.add_stage(_cell(r, mapping.get("stage")))
        trial = parse_number(_cell(r, mapping.get("trial")))
        video = _cell(r, mapping.get("video"))
        if video and video_dir is not None:
            if re.match(r"^[A-Za-z]:[\\/]", video) or "\\" in video:  # a Windows path from ANY-maze
                video = video.replace("\\", "/").rsplit("/", 1)[-1]
            if not Path(video).is_absolute():
                video = str(Path(video_dir) / video)
        app = _cell(r, mapping.get("apparatus"))
        app = app if app in [x.name for x in project.apparatus] else ""
        t = project.add_test(video, aid, app, stage=stage, trial=int(trial) if math.isfinite(trial) else 1)
        dur = parse_number(_cell(r, mapping.get("duration")))
        if math.isfinite(dur) and dur > 0:
            t.duration_s = dur
        new.append(t)
    return new


def import_track(header, rows, mapping: dict, scale: float = 1.0, offset=(0.0, 0.0), flip_y: bool = False,
                 height: float = 0.0, fps: float | None = None) -> Track:
    """A Track from a table of positions.

    Positions are converted to video pixels as ``px = value * scale + offset`` (e.g. ANY-maze positions in metres
    relative to the apparatus: scale = pixels per metre). ``flip_y`` mirrors y (``height - y``) for software whose y
    axis points up.
    """
    def col(role):
        i = mapping.get(role)
        if i is None:
            return None
        return np.array([parse_number(_cell(r, i)) for r in rows], float)

    t = col("t")
    if t is None:
        if not fps:
            raise ValueError("the table has no time column: give the frame rate")
        t = np.arange(len(rows)) / fps
    good = np.isfinite(t)
    vals = {}
    for role in ("x", "y", "hx", "hy", "tx", "ty"):
        v = col(role)
        if v is None:
            v = np.full(len(rows), np.nan)
        else:
            v = v * scale + (offset[0] if role.endswith("x") else offset[1])
            if flip_y and role.endswith("y"):
                v = height - v
        vals[role] = v[good]
    t = t[good]
    if vals["x"].size == 0 or not np.isfinite(vals["x"]).any():
        raise ValueError("no positions found: map the centre X / Y columns")
    t = t - t[0]
    dt = np.diff(t)
    est = 1.0 / float(np.median(dt[dt > 0])) if (dt > 0).any() else (fps or 25.0)
    detected = np.isfinite(vals["x"]) & np.isfinite(vals["y"])
    tr = Track(t=t, **vals, detected=detected, fps=fps or est)
    tr.meta["source"] = "imported table"
    return tr


def dlc_bodyparts(path) -> list[str]:
    """Body parts of a DeepLabCut CSV in column order ([] if the file is not one)."""
    with open(path, newline="") as f:
        for i, row in enumerate(csv.reader(f)):
            if row and row[0].strip().lower() == "bodyparts":
                return list(dict.fromkeys(row[1:]))
            if i > 5:
                break
    return []


def trim_to_test(track: Track, start_s: float, duration_s: float) -> Track:
    """The test period (from start_s, duration_s seconds; 0 = to the end) of a track covering the whole video, in
    test time."""
    if start_s > 0 or duration_s:
        track = track.slice_time(start_s, start_s + duration_s if duration_s else math.inf)
        track.t = track.t - start_s
    track.meta["video_start_s"] = start_s
    return track
