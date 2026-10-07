"""Visualisation, statistics and data-transfer features: per-frame charts, plots, statistics, exports, video."""

import csv
from collections import OrderedDict
from dataclasses import asdict

import numpy as np
import pytest
from scipy import stats as sps

from manymaze.core import charts, plots, stats as st, templates
from manymaze.core.demo import create_demo_project
from manymaze.core.export import (export_raw_data, export_results, export_xml, html_report, read_experiment_xml,
                                  table_text, write_table)
from manymaze.core.measures import AnalysisSettings, analyse
from manymaze.core.project import Behaviour
from manymaze.core.track import Track


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    return create_demo_project(tmp_path_factory.mktemp("demo") / "d.mmaze", n_per_group=1, seconds=5)


def _circle_track(n=750, fps=25.0, head=True):
    t = np.arange(n) / fps
    x = 200 + 100 * np.cos(t / 3)
    y = 200 + 100 * np.sin(t / 3)
    kw = dict(hx=x + 6, hy=y, tx=x - 6, ty=y, angle=np.zeros(n)) if head else {}
    return Track(t=t, x=x, y=y, area=np.full(n, 300.0), motion=np.where(t % 10 < 3, 1.0, 40.0), fps=fps, **kw)


@pytest.fixture(scope="module")
def app():
    a = templates.build("open_field", 50, 50, 300, 300, size_cm=40)
    a.frame_size = (400, 400)
    return a


# ---------------------------------------------------------------- charts
def test_parameters_and_values(app):
    tr = _circle_track()
    beh = [Behaviour("Rearing", kind="state"), Behaviour("Poop", kind="point")]
    params = charts.parameters(app, tr, beh)
    names = [p.name for p in params]
    assert len(params) >= 40 and len(set(names)) == len(names)
    general = [p for p in params if p.group not in ("Zones", "Points", "Lines", "Behaviours", "Social")]
    assert len(general) >= 30
    events = [{"behaviour": "Rearing", "t": 3.0, "t_end": 5.0}, {"behaviour": "Poop", "t": 7.0, "t_end": None}]
    d = charts.compute(tr, app, None, AnalysisSettings(), events, beh)
    assert set(d) == set(names) and all(len(v) == len(tr) for v in d.values())
    res = analyse(tr, app, AnalysisSettings(), events, beh)
    assert d["Distance travelled"][-1] == pytest.approx(res["Total distance (cm)"], rel=1e-3)
    assert d["Centre: time in zone"][-1] == pytest.approx(res["Centre: time (s)"], abs=0.05)
    assert d["Centre: entries"][-1] == res["Centre: entries"]
    assert d["Time freezing"][-1] == pytest.approx(res["Time freezing (s)"], abs=0.05)
    assert d["Rearing: active"][(tr.t >= 3) & (tr.t < 5)].all() and d["Rearing: active"][tr.t >= 5].sum() == 0
    assert d["Poop: count"][tr.t < 7].max() == 0 and d["Poop: count"][-1] == 1
    inside = d["Centre: in zone"] > 0
    assert (d["Centre: distance to zone"][inside] == 0).all() and (d["Centre: distance to zone"][~inside] > 0).all()
    assert np.nanmax(np.abs(d["Distance from arena centre"] - 100 / 7.5)) < 1e-6
    assert d["Speed"][5:-5] == pytest.approx(np.full(len(tr) - 10, 100 / 3 / 7.5), rel=0.02)
    assert np.allclose(d["Body length"], 12 / 7.5)
    with pytest.raises(KeyError):
        charts.compute(tr, app, ["Nope"])
    # without head data the head parameters are not offered
    assert "Head X position" not in [p.name for p in charts.parameters(app, _circle_track(head=False))]


def test_measure_peaks_and_figure(app):
    t = np.arange(0, 20, 0.04)
    v = np.sin(2 * np.pi * t / 5)
    m = charts.measure_interval(t, v, 0, 5)
    assert m["max"] == pytest.approx(1, abs=1e-3) and m["t_max"] == pytest.approx(1.25, abs=0.05)
    assert m["mean"] == pytest.approx(0, abs=0.02) and m["duration"] == 5
    assert len(charts.find_peaks(t, v)) == 4 and len(charts.find_peaks(t, v, valleys=True)) == 4
    tr = _circle_track()
    beh = [Behaviour("Rearing", kind="state")]
    fig = plots.chart_figure(tr, app, ["Speed", "Freezing", "Centre: entries"], events=[
        {"behaviour": "Rearing", "t": 3.0, "t_end": 5.0}], behaviours=beh, bands=["Centre", "Corners"])
    assert len(fig.axes) == 4  # 3 parameters + event strip
    assert fig.axes[0].get_legend() is not None  # zone band legend
    cols, arr = charts.per_frame_table(tr, app, names=["Speed", "X position"])
    assert cols == ["Time (s)", "Speed (cm/s)", "X position (cm)"] and arr.shape == (len(tr), 3)
    assert charts.state_mask(tr, app, "Freezing").sum() > 0


# ---------------------------------------------------------------- plots
def test_alignment_transforms(app):
    b = (0, 0, 100, 50)
    x, y = plots.align_xy([100], [25], b, "rot90")  # right-middle → bottom-middle (clockwise on screen)
    assert (x[0], y[0]) == pytest.approx((50, 50))
    x, y = plots.align_xy([10], [5], b, "flipx")
    assert (x[0], y[0]) == pytest.approx((90, 5))
    px, py = np.array([3.0, 70.0]), np.array([4.0, 33.0])
    x, y = px, py
    for _ in range(4):
        x, y = plots.align_xy(x, y, b, "rot90")
    assert np.allclose(x, px) and np.allclose(y, py)
    x, y = plots.align_xy([0], [0], b, "none", (200, 200, 400, 300))
    assert (x[0], y[0]) == pytest.approx((200, 200))
    with pytest.raises(ValueError):
        plots.align_xy([0], [0], b, "spin")
    tr = _circle_track()
    other = templates.build("open_field", 0, 0, 600, 600, size_cm=40)
    al = plots.align_track(tr, app, "rot180", other)
    assert al.x[0] == pytest.approx(600 - (tr.x[0] - 50) * 2, abs=1e-6)
    assert plots.align_track(tr, app, "none") is tr


def test_heatmaps_and_track_plots(app):
    tr = _circle_track()
    H, ext = plots.occupancy([tr], app, norm="percent")
    assert H.sum() == pytest.approx(100, rel=0.02)
    Ht, _ = plots.occupancy([tr], app)
    assert Ht.sum() == pytest.approx(tr.frame_durations().sum(), rel=0.02)
    mask = tr.t < 10
    Hm, _ = plots.occupancy([tr], app, masks=[mask])
    Hr, _ = plots.occupancy([tr], app, t_range=(0, 10))
    assert np.allclose(Hm, Hr) and Hm.sum() < Ht.sum() / 2
    assert plots.occupancy([tr], app, norm="relative")[0].max() == pytest.approx(1)
    fig = plots.heatmap([tr], app, norm="percent", masks=[mask])
    assert "% of time" in fig.axes[-1].get_ylabel()
    g = plots.group_heatmap_figure([("A", [tr], app), ("B", [plots.align_track(tr, app, "flipy")], app)],
                                   norm="percent")
    assert len([a for a in g.axes if a.images]) == 2
    mk = plots.behaviour_markers(tr, app, AnalysisSettings(), [{"behaviour": "Poop", "t": 2.0, "t_end": None}],
                                 [Behaviour("Poop", kind="point")])
    assert {m["label"] for m in mk} == {"Freezing", "Poop"}
    for color_by in ("none", "time", "speed", "Distance from wall"):
        fig = plots.track_plot(tr, app, color_by=color_by, colorbar=True, markers=mk, part="head")
        assert len(fig.axes) == (1 if color_by == "none" else 2)
    seg = plots.segmented_track_plot(tr, app, [("0-10 s", 0, 10), ("10-20 s", 10, 20), ("20-30 s", 20, 30)],
                                     color_by="speed", markers=mk)
    assert len([a for a in seg.axes if a.get_title()]) == 3


def test_group_graphs():
    rng = np.random.default_rng(0)
    rows = [{"Group": g, "Stage": s, "Sex": x, "v": float(rng.normal(i))} for i, g in enumerate("AB") for s in
            ("D1", "D2", "D3") for x in ("F", "M") for _ in range(3)]
    for kind in ("bar", "line", "point", "box", "violin"):
        fig = plots.factor_plot(rows, "v", ["Stage", "Group", "Sex"], kind=kind, error="ci")
        assert len(fig.axes) == 2  # one panel per sex
    gv = st.group_values(rows, "v")
    for kind in ("bar", "box", "violin", "point"):
        assert plots.group_plot(gv, "v", kind=kind, error="sd", points=False).axes
    fig = plots.time_course(rows, "v", x="Stage", by="Group", error="ci", points=True)
    assert len(fig.axes[0].get_legend().get_texts()) == 2
    x = rng.normal(size=20)
    fig = plots.scatter_plot(x, 2 * x, ["g"] * 20, "x", "y", res=st.correlation(x, 2 * x))
    assert fig.axes[0].collections


# ---------------------------------------------------------------- statistics
def test_statistics_catalogue_and_compat():
    assert len(st.TESTS) >= 30
    rng = np.random.default_rng(3)
    g = OrderedDict(A=rng.normal(0, 1, 10), B=rng.normal(1.5, 1, 10))
    r = st.compare_groups(g)
    assert r["test"] == "Welch's t-test" and r["p"] == pytest.approx(sps.ttest_ind(g["A"], g["B"],
                                                                                    equal_var=False).pvalue)
    assert st.compare_groups(g, parametric=False)["test"] == "Mann-Whitney U"
    assert st.compare_groups(g, paired=True)["test"] == "Paired t-test"
    g3 = OrderedDict(A=g["A"], B=g["B"], C=rng.normal(3, 1, 10))
    r = st.compare_groups(g3)
    assert r["test"] == "One-way ANOVA" and len(r["posthoc"]) == 3 and r["posthoc"][0]["test"] == "Tukey HSD"
    assert st.compare_groups(g3, parametric=False)["posthoc"][0]["test"] == "Mann-Whitney (Bonferroni)"
    for m in ("welch_anova", "alexander_govern", "kruskal", "median", "rm_anova", "friedman", "student", "ks"):
        res = st.compare_groups(g3, method=m)
        assert 0 <= res["p"] <= 1, m
    for ph in ("tukey", "bonferroni", "holm", "sidak", "fdr", "dunnett", "games_howell", "dunn", "duncan", "lsd",
               "scheffe", "snk"):
        res = st.compare_groups(g3, posthoc_method=ph, control="B")
        assert res["posthoc"] and all(0 <= x["p"] <= 1 for x in res["posthoc"]), ph
    assert len(st.compare_groups(g3, posthoc_method="dunnett", control="B")["posthoc"]) == 2
    assert st.two_way_anova([], "v")["error"]


# sweet potato yield by virus (R package agricolae, data(sweetpotato)): MSE = 22.48917 on 8 df
SWEETPOTATO = OrderedDict(cc=[28.5, 21.7, 23.0], fc=[14.9, 10.6, 13.1], ff=[41.8, 39.2, 28.0],
                          oo=[38.2, 40.4, 32.1])


def test_range_posthoc_tests_against_published_values():
    mse, df = st.anova_error_term([np.array(v) for v in SWEETPOTATO.values()])
    assert mse == pytest.approx(22.48917, abs=1e-5) and df == 8
    s = (mse / 3) ** 0.5
    # least significant ranges printed by agricolae's duncan.test / SNK.test / LSD.test / HSD.test (alpha = 0.05)
    duncan = [st.range_test_critical(r, df, "duncan") * s for r in (2, 3, 4)]
    snk = [st.range_test_critical(r, df, "snk") * s for r in (2, 3, 4)]
    assert duncan == pytest.approx([8.928965, 9.304825, 9.514910], abs=2e-5)
    assert snk == pytest.approx([8.928965, 11.064170, 12.399670], abs=2e-5)
    # Duncan's table of significant studentized ranges (Harter 1960), 8 error df: 3.26, 3.40, 3.47
    assert [st.range_test_critical(r, 8, "duncan") for r in (2, 3, 4)] == pytest.approx([3.261, 3.399, 3.475],
                                                                                         abs=1e-3)
    # a difference equal to the least significant range has p = alpha
    for r, crit in zip((2, 3, 4), duncan):
        assert st.range_test_p(crit / s, r, df, "duncan") == pytest.approx(0.05, abs=1e-6)
    for r, crit in zip((2, 3, 4), snk):
        assert st.range_test_p(crit / s, r, df, "snk") == pytest.approx(0.05, abs=1e-6)

    def sig(method, **kw):
        res = st.posthoc(SWEETPOTATO, method, **kw)
        return {frozenset((x["a"], x["b"])) for x in res if x["p"] < 0.05}, {(x["a"], x["b"]): x for x in res}

    pairs = {frozenset(p) for p in [("cc", "fc"), ("cc", "ff"), ("cc", "oo"), ("fc", "ff"), ("fc", "oo")]}
    # agricolae groups: Duncan, SNK and LSD: oo a, ff a, cc b, fc c; HSD: oo a, ff ab, cc bc, fc c
    for m in ("duncan", "snk", "lsd"):
        assert sig(m)[0] == pairs, m
    assert sig("tukey")[0] == {frozenset(p) for p in [("cc", "oo"), ("fc", "ff"), ("fc", "oo")]}
    # Scheffé's critical difference sqrt((k-1) F(0.95; 3, 8)) * sqrt(2 MSE / n) = 13.52: only fc differs
    sch, by = sig("scheffe")
    assert sch == {frozenset(p) for p in [("fc", "ff"), ("fc", "oo")]}
    f_crit = sps.f.ppf(0.95, 3, 8)
    assert f_crit == pytest.approx(4.066181, abs=1e-6)
    msd = (3 * f_crit) ** 0.5 * (2 * mse / 3) ** 0.5
    assert msd == pytest.approx(13.5237, abs=1e-3)
    # exact values: LSD = pooled t-test; Scheffé F = t² / (k - 1)
    t = 11.533333 / (2 * mse / 3) ** 0.5
    _, lsd = sig("lsd")
    assert lsd[("cc", "fc")]["p"] == pytest.approx(2 * sps.t.sf(t, 8), rel=1e-5)
    assert by[("cc", "fc")]["p"] == pytest.approx(sps.f.sf(t ** 2 / 3, 3, 8), rel=1e-5)
    # SNK over the whole range = Tukey HSD; Duncan over 2 adjacent means = LSD
    _, snk_res = sig("snk")
    tk = sps.tukey_hsd(*[np.array(v) for v in SWEETPOTATO.values()])
    assert snk_res[("fc", "oo")]["p"] == pytest.approx(tk.pvalue[1, 3], rel=1e-4)
    _, dun = sig("duncan")
    assert dun[("cc", "fc")]["p"] == pytest.approx(lsd[("cc", "fc")]["p"], rel=1e-6)
    # step-down: no pair is more significant than a wider range that contains it
    assert snk_res[("fc", "ff")]["p"] >= snk_res[("fc", "oo")]["p"]
    assert snk_res[("fc", "ff")]["p_unadjusted"] < snk_res[("fc", "ff")]["p"]
    # ordering of conservativeness for every pair
    for key in lsd:
        assert lsd[key]["p"] <= dun[key]["p"] + 1e-12 <= snk_res[key]["p"] + 1e-9 <= by[key]["p"] + 1e-9


def test_range_posthoc_tests_two_groups_and_designs():
    rng = np.random.default_rng(11)
    a, b = rng.normal(0, 1, 7), rng.normal(1.2, 1, 9)
    student = sps.ttest_ind(a, b).pvalue  # with two groups every test reduces to Student's t-test
    for m in ("lsd", "scheffe", "snk", "duncan"):
        assert st.posthoc(OrderedDict(a=a, b=b), m)[0]["p"] == pytest.approx(student, rel=1e-4), m
    # paired: the error term of the repeated-measures ANOVA; with 2 conditions LSD = paired t-test
    c = a + rng.normal(0.4, 0.3, 7)
    assert st.posthoc(OrderedDict(a=a, c=c), "lsd", paired=True)[0]["p"] == pytest.approx(
        sps.ttest_rel(a, c).pvalue, rel=1e-6)
    # through compare_groups (one-way ANOVA, unequal n)
    g3 = OrderedDict(A=a, B=b, C=rng.normal(3, 1, 6))
    for m, name in (("duncan", "Duncan"), ("lsd", "Fisher's LSD"), ("scheffe", "Scheffé"),
                    ("snk", "Student-Newman-Keuls")):
        res = st.compare_groups(g3, method="anova", posthoc_method=m)
        assert len(res["posthoc"]) == 3 and all(x["test"] == name and 0 <= x["p"] <= 1 for x in res["posthoc"])
    # degenerate data do not raise
    assert st.posthoc(OrderedDict(a=[1.0, 1.0], b=[1.0, 1.0], c=[2.0, 2.0]), "snk")[0]["p"] == 1.0


def test_statistics_against_references():
    rng = np.random.default_rng(5)
    a, b = rng.normal(0, 1, 12), rng.normal(1, 2, 9)
    # Welch ANOVA with 2 groups == Welch t-test (F = t²)
    w = st.compare_groups(OrderedDict(a=a, b=b), method="welch_anova")
    t = sps.ttest_ind(a, b, equal_var=False)
    assert w["test"] == "Welch's t-test" or w["p"] == pytest.approx(t.pvalue)
    F, df, p = st._welch_anova([a, b])
    assert F == pytest.approx(t.statistic ** 2) and p == pytest.approx(t.pvalue)
    # Games-Howell for 2 groups == Welch t-test
    gh = st.posthoc(OrderedDict(a=a, b=b), "games_howell")[0]
    assert gh["p"] == pytest.approx(t.pvalue, rel=1e-3)
    # RM ANOVA with 2 conditions == paired t² ; GG epsilon = 1
    c = a + rng.normal(0.5, 0.5, 12)
    rm = st.rm_anova_matrix(np.column_stack([a, c]))
    pt = sps.ttest_rel(a, c)
    assert rm["F"] == pytest.approx(pt.statistic ** 2) and rm["p"] == pytest.approx(pt.pvalue)
    assert rm["epsilon_gg"] == pytest.approx(1.0)
    # p-value adjustment
    assert st.p_adjust([0.01, 0.04, 0.03], "holm") == pytest.approx([0.03, 0.06, 0.06])
    assert st.p_adjust([0.01, 0.04, 0.03], "fdr") == pytest.approx([0.03, 0.04, 0.04])
    assert st.p_adjust([0.01, 0.02], "bonferroni") == pytest.approx([0.02, 0.04])
    # one-sample
    o = st.one_sample(a, 0.3)
    assert o["p"] == pytest.approx(sps.ttest_1samp(a, 0.3).pvalue)
    assert st.one_sample(a, 0.3, parametric=False)["test"] == "Wilcoxon signed-rank vs value"
    # regression / correlation
    x = np.linspace(0, 10, 30)
    y = 2 * x + 1 + rng.normal(0, 0.5, 30)
    reg = st.regression(x, y)
    lr = sps.linregress(x, y)
    half = sps.t.ppf(0.975, 28) * lr.stderr
    assert reg["slope_ci"] == pytest.approx((lr.slope - half, lr.slope + half))
    assert reg["r2"] > 0.95 and len(reg["band_x"]) == 60 and (reg["band_lo"] < reg["band_y"]).all()
    assert st.correlation(x, y, "kendall")["r"] > 0.8
    assert st.correlation(x, y)["ci95"][0] < st.correlation(x, y)["r"]
    # assumptions
    assert set(st.normality(rng.normal(size=20))) == {"Shapiro-Wilk", "D'Agostino-Pearson"}
    assert set(st.variance_tests({"a": a, "b": b})) == {"Levene", "Brown-Forsythe", "Bartlett", "Fligner-Killeen"}


def test_two_factor_designs():
    rng = np.random.default_rng(2)
    rows = []
    for grp, slope in (("Ctl", 1.0), ("Drug", 2.5)):
        for i in range(6):
            base = rng.normal(0, 1)
            for j, stage in enumerate(("D1", "D2", "D3")):
                rows.append({"Group": grp, "Animal": f"{grp}{i}", "Stage": stage, "v": base + slope * j +
                             rng.normal(0, 0.3), "Strategy": "Direct" if (grp == "Drug") ^ (i < 1) else "Random"})
    mx = st.mixed_anova(rows, "v", "Group", "Stage")
    eff = {e["effect"]: e for e in mx["effects"]}
    assert set(eff) == {"Group", "Stage", "Group × Stage"} and eff["Group × Stage"]["p"] < 0.001
    assert eff["Group"]["df_error"] == 10 and eff["Stage"]["df_error"] == 20 and "p_gg" in eff["Stage"]
    # between effect == one-way ANOVA on the animals' means
    means = {}
    for r in rows:
        means.setdefault((r["Group"], r["Animal"]), []).append(r["v"])
    ctl = [np.mean(v) for (g, _), v in means.items() if g == "Ctl"]
    drug = [np.mean(v) for (g, _), v in means.items() if g == "Drug"]
    assert eff["Group"]["p"] == pytest.approx(sps.f_oneway(ctl, drug).pvalue)
    rm = st.rm_anova(rows, "v", "Stage")
    assert rm["df"] == (2, 22) and rm["p"] < 0.001
    for fn in (st.two_way_anova, st.scheirer_ray_hare, st.art_anova):
        res = fn(rows, "v", "Group", "Stage")
        assert [e["effect"] for e in res["effects"]] == ["Group", "Stage", "Group × Stage"]
        assert res["effects"][1]["p"] < 0.001
    assert "Greenhouse" in st.anova_text(mx, "v")
    # categorical
    rl, cl, T = st.contingency_table(rows, "Group", "Strategy")
    res = st.categorical_test(T, rl, cl)
    assert T.sum() == len(rows) and res["p"] < 0.05 and "fisher" in res and 0 <= res["effect_size"] <= 1
    assert res["fisher"]["p"] == pytest.approx(sps.fisher_exact(T)[1])
    assert st.chi_square_gof([10, 10, 10])["p"] == pytest.approx(1.0)
    # descriptive tables up to three levels
    d = st.describe_by(rows, "v", ["Group", "Stage"], {"Stage": ["D3", "D2", "D1"]})
    assert len(d) == 6 and [x["Stage"] for x in d[:3]] == ["D3", "D2", "D1"] and d[0]["n"] == 6
    assert len(st.describe_by(rows, "v", ["Group", "Stage", "Strategy"])) == 12


# ---------------------------------------------------------------- export
def test_xml_and_raw_exports(demo, tmp_path):
    t = demo.tests[0]
    t.io_events = [{"t": 1.5, "device": "virtual", "channel": "led", "kind": "output", "value": 1}]
    t.variables["heatmap_transform"] = "rot90"
    path = export_xml(demo, tmp_path / "e.xml")
    data = read_experiment_xml(path)
    assert data["name"] == demo.name and len(data["tests"]) == len(demo.tests)
    t0 = data["tests"][0]
    assert t0["io_events"][0]["channel"] == "led" and t0["events"][0]["behaviour"] == "Rearing"
    tr = demo.load_tracks(t)[0]
    cols = t0["tracks"][0]["columns"]
    assert set(cols) >= {"t", "x", "y", "detected"} and np.allclose(cols["x"], tr.x, atol=1e-3, equal_nan=True)
    whole = next(r for r in t0["results"] if r["period"] == "Whole test")["values"]
    ref = demo.analyse_test(t)[0]
    assert whole["Total distance (cm)"] == pytest.approx(ref["Total distance (cm)"])
    text = path.read_text()
    assert "<apparatus " in text and "<zone-group" in text and 'name="heatmap_transform"' in text
    # results without tracks, via export_results
    p2 = export_results(demo, tmp_path / "r.xml")
    assert p2.exists()
    # per-test raw data with derived parameters
    files = export_raw_data(demo, tmp_path / "raw")
    assert len(files) == len(demo.tests)
    with open(files[0]) as f:
        lines = [ln for ln in f if not ln.startswith("#")]
    head = next(csv.reader([lines[0]]))
    assert head[:3] == ["Time (s)", "raw x", "raw y"] and "Speed (cm/s)" in head and "Centre: in zone" in head
    assert len(lines) - 1 == len(tr)
    files = export_raw_data(demo, tmp_path / "raw2", parameters=["Speed"], delimiter="\t")
    assert files[0].suffix == ".tsv" and open(files[0]).read().count("Speed (cm/s)") == 1


def test_tables_and_report(demo, tmp_path):
    rows = demo.results()
    cols = ["Test", "Animal", "Total distance (cm)"]
    for ext in (".csv", ".tsv", ".xlsx"):
        assert write_table(rows, tmp_path / f"t{ext}", cols).exists()
    assert (tmp_path / "t.tsv").read_text().splitlines()[0] == "Test\tAnimal\tTotal distance (cm)"
    txt = table_text(rows, cols)
    assert txt.count("\n") == len(rows) + 1
    export_results(demo, tmp_path / "r.tsv")
    assert "\t" in (tmp_path / "r.tsv").read_text().splitlines()[0]
    from openpyxl import load_workbook

    wb = load_workbook(export_results(demo, tmp_path / "r.xlsx"))
    assert wb.sheetnames == ["Results", "Zone visits", "Animals", "Tests", "Settings"]
    assert wb["Results"].max_row == len(rows) + 1
    rep = html_report(demo, tmp_path / "r.html", stats_measures=["Total distance (cm)"], heatmap_norm="fixed",
                      chart_parameters=["Speed", "Centre: in zone"], color_by="speed")
    text = rep.read_text()
    assert text.count("<img") >= 3 * len(demo.tests) + 1 and "Group occupancy" in text


def test_group_heatmap(demo):
    by_group = {}
    for t in demo.tests:
        by_group.setdefault(demo.get_animal(t.animal_id).group, []).append(t)
    seen = []
    fig = plots.group_heatmap(demo, dict(reversed(list(by_group.items()))), heat_of="Freezing", progress=seen.append)
    titles = [a.get_title() for a in fig.axes if a.images]
    assert [t.split(" (")[0] for t in titles] == [g.name for g in demo.groups] and seen[-1] == 1.0
    fig = plots.group_heatmap(demo, by_group, period="no such period")
    assert all(a.get_title().endswith("(n = 0)") for a in fig.axes if a.images)


# ---------------------------------------------------------------- video
def test_video_export(demo, tmp_path):
    from manymaze.core.video import VideoSource
    from manymaze.core.videoexport import OverlayOptions, OverlayRenderer, export_video

    t = demo.tests[0]
    seen = []
    out = export_video(demo, t, tmp_path / "o.mp4", OverlayOptions(trail_s=-1, t_start=1.0, t_end=3.0),
                       progress=seen.append)
    with VideoSource(str(out)) as v:
        assert v.width == 400 and abs(v.frame_count - 51) <= 1
        f = v.frame_at(25)
    assert seen and seen[-1] == 1.0 and f is not None
    out = export_video(demo, t, tmp_path / "s.avi", OverlayOptions(scale=0.5, speed=2.0, trail_color="time"))
    with VideoSource(str(out)) as v:
        assert v.width == 200 and abs(v.frame_count - 63) <= 2
    assert export_video(demo, t, tmp_path / "c.mp4", should_stop=lambda: True) is None
    assert not (tmp_path / "c.mp4").exists()
    # the renderer draws a behaviour label while a scored state is active
    tr = demo.load_tracks(t)[0]
    r = OverlayRenderer([tr], demo.apparatus[0], OverlayOptions(timestamp=False, info=False, zones=False,
                                                                trail_s=0, body_points=False, freezing=False),
                        [{"behaviour": "Rearing", "t": 1.0, "t_end": 2.0}], [Behaviour("Rearing", kind="state")])
    blank = np.zeros((400, 400, 3), np.uint8)
    assert r.render(blank, 1.5).sum() > 0 and r.render(blank, 3.0).sum() == 0
