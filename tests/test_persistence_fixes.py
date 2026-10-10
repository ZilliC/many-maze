"""Regression tests for the persistence / import / export / CLI audit (docs/AUDIT-2026-10-08.md, "Persistence,
import/export, CLI")."""

import json
import math
import os
import re
import shutil
import stat
import subprocess
import sys
import zipfile

import numpy as np
import pytest

from manymaze import cli
from manymaze.core import autosave, explock, export, workflow
from manymaze.core.anymaze import day_first, import_anymaze_xml, read_anymaze_xml
from manymaze.core.apparatus import Apparatus
from manymaze.core.archive import archive_project
from manymaze.core.atomicfile import atomic_write
from manymaze.core.geometry import rect
from manymaze.core.importers import decimal_comma, dlc_bodyparts, import_track, parse_number, read_table
from manymaze.core.ioconfig import is_secret, new_device
from manymaze.core.project import (ERROR_COLUMN, PROJECT_FILE, SECRETS_FILE, Animal, Group, Project, relink_videos,
                                   same_folder)
from manymaze.core.track import COLUMNS, Track


def _track(n=50, fps=10.0, **meta):
    t = np.arange(n) / fps
    tr = Track(t=t, x=100 + 10 * np.sin(t), y=100 + 10 * np.cos(t), fps=fps)
    tr.meta.update(meta)
    return tr


def _project(tmp_path, name="e.mmaze", n_tests=2) -> Project:
    p = Project(name="exp")
    p.apparatus.append(Apparatus(name="OF", arena=rect(50, 50, 100, 100)))
    p.groups = [Group("Saline"), Group("Drug")]
    p.animals = [Animal("A1", "Saline"), Animal("A2", "Drug")]
    p.save(tmp_path / name)
    for i in range(n_tests):
        t = p.add_test("", f"A{i + 1}", "OF")
        p.save_tracks(t, [_track()])
    p.save()
    return p


# ---------------------------------------------------------------- decimal separator per file
def test_comma_decimals_with_three_places_are_not_thousands():
    rows = [["0,125", "0,040", "1,5"], ["0,250", "2,125", "1,25"]]
    header = ["Time", "Centre X", "Centre Y"]
    tr = import_track(header, rows, {"t": 0, "x": 1, "y": 2})
    np.testing.assert_allclose(tr.x, [0.04, 2.125])
    np.testing.assert_allclose(tr.t, [0.0, 0.125])  # t starts at 0: 0.125 and 0.25
    assert decimal_comma(["0,125", "12,500"]) is True
    assert parse_number("0,125") == 0.125 and parse_number("0,040") == 0.04  # a thousands group never starts at 0
    assert parse_number("1.234,5") == 1234.5
    assert parse_number("12,500", decimal_comma=True) == 12.5


def test_thousands_separators_kept_when_the_file_uses_them():
    rows = [["1,234,567", "2.5"], ["12", "1,234.5"], ["1,234", "3"]]
    assert decimal_comma(c for r in rows for c in r) is False
    tr = import_track(["Time", "X", "Y"], [[str(i), r[0], r[1]] for i, r in enumerate(rows)], {"t": 0, "x": 1, "y": 2})
    np.testing.assert_allclose(tr.x, [1234567, 12, 1234])
    np.testing.assert_allclose(tr.y, [2.5, 1234.5, 3])
    assert parse_number("1,234,567") == 1234567 and parse_number("1,234") == 1234  # alone: thousands
    assert decimal_comma(["1,234", "12"]) is None  # nothing tells


def test_semicolon_csv_with_comma_decimals(tmp_path):
    f = tmp_path / "track.csv"
    f.write_text("Time;Centre position X;Centre position Y\n0,000;0,125;1,000\n0,040;0,250;1,500\n", encoding="utf-8")
    h, rows = read_table(f)
    tr = import_track(h, rows, {"t": 0, "x": 1, "y": 2})
    np.testing.assert_allclose(tr.x, [0.125, 0.25])


# ---------------------------------------------------------------- track CSV encoding and meta
def test_track_csv_is_utf8_and_reads_cp1252(tmp_path):
    tr = _track(video="/Vidéos/実験/test 1.mp4", note="line 1\nline 2", pad=" spaced ")
    f = tmp_path / "t.csv"
    tr.to_csv(f)
    raw = f.read_bytes()
    raw.decode("utf-8")  # strictly UTF-8
    assert b'# note:json="line 1\\nline 2"' in raw  # one header line: the data still parses
    back = Track.from_csv(f)
    assert back.meta["video"] == "/Vidéos/実験/test 1.mp4"
    assert back.meta["note"] == "line 1\nline 2" and back.meta["pad"] == " spaced "
    assert len(back) == len(tr)
    # a file written by an older version on Windows (cp1252)
    old = tmp_path / "old.csv"
    old.write_bytes("# fps=10.0\n# video=C:\\Vidéos\\a.mp4\nt,x,y,detected\n0,1,2,1\n".encode("cp1252"))
    assert Track.from_csv(old).meta["video"] == "C:\\Vidéos\\a.mp4"


def test_dlc_bodyparts_utf8(tmp_path):
    f = tmp_path / "dlc.csv"
    f.write_text("scorer,DLC,DLC\nbodyparts,museau,queue\ncoords,x,y\n0,1,2\n", encoding="utf-8")
    assert dlc_bodyparts(f) == ["museau", "queue"]
    f.write_bytes("scorer,DLC,DLC\nbodyparts,tête,queue\ncoords,x,y\n".encode("cp1252"))
    assert dlc_bodyparts(f) == ["tête", "queue"]


# ---------------------------------------------------------------- secrets
def _with_alerts(p: Project) -> Project:
    d = new_device("notify")
    d.update(smtp_host="smtp.lab", smtp_user="lab", smtp_password="hunter2", twilio_token="tok-123")
    p.io_devices = [d]
    return p


def test_secrets_kept_out_of_project_json(tmp_path):
    assert is_secret("smtp_password") and is_secret("twilio_token") and is_secret("api_key")
    assert not is_secret("smtp_user")
    p = _with_alerts(_project(tmp_path))
    p.save()
    text = (p.path / PROJECT_FILE).read_text(encoding="utf-8")
    assert "hunter2" not in text and "tok-123" not in text and "smtp.lab" in text
    sec = p.path / SECRETS_FILE
    assert "hunter2" in sec.read_text(encoding="utf-8")
    if os.name != "nt":
        assert stat.S_IMODE(sec.stat().st_mode) == 0o600
    q = Project.load(p.path)
    assert q.io_devices[0]["smtp_password"] == "hunter2" and q.io_devices[0]["twilio_token"] == "tok-123"
    # a project.json of an older version with the password in it: moved out at the next save
    d = json.loads(text)
    d["io_devices"][0]["smtp_password"] = "old-pass"
    (p.path / PROJECT_FILE).write_text(json.dumps(d), encoding="utf-8")
    sec.unlink()
    r = Project.load(p.path)
    assert r.io_devices[0]["smtp_password"] == "old-pass"
    r.save()
    assert "old-pass" not in (p.path / PROJECT_FILE).read_text(encoding="utf-8")
    assert "old-pass" in sec.read_text(encoding="utf-8")
    # no secrets left: no file
    r.io_devices[0]["smtp_password"] = r.io_devices[0]["twilio_token"] = ""
    r.save()
    assert not sec.exists()


def test_secrets_not_in_report_archive_or_protocol_copy(tmp_path):
    p = _with_alerts(_project(tmp_path))
    p.save()
    html = export.protocol_report(p, tmp_path / "protocol.html").read_text(encoding="utf-8")
    assert "smtp.lab" in html and "hunter2" not in html and "tok-123" not in html
    z = archive_project(p, tmp_path / "a.zip", include_videos=False)
    with zipfile.ZipFile(z) as zf:
        assert not any(n.endswith(SECRETS_FILE) or n.endswith(explock.LOCK_FILE) for n in zf.namelist())
        assert all(b"hunter2" not in zf.read(n) for n in zf.namelist())
    dst = workflow.copy_protocol(p, Project(name="new"))
    assert dst.io_devices[0]["smtp_host"] == "smtp.lab" and "smtp_password" not in dst.io_devices[0]
    assert p.io_devices[0]["smtp_password"] == "hunter2"  # the source keeps its own


# ---------------------------------------------------------------- XML control characters
def test_xml_export_strips_illegal_characters(tmp_path):
    p = _project(tmp_path)
    p.description = "bad \x01 char"
    p.tests[0].notes = "vertical\x0btab"
    p.animals[0].notes = "bell\x07"
    p.animals[0].fields["Note"] = "x\x1fy"
    out = export.export_xml(p, tmp_path / "e.xml")
    d = export.read_experiment_xml(out)
    assert len(d["tests"]) == 2 and d["animals"][0]["notes"] == "bell"


# ---------------------------------------------------------------- videos outside the experiment folder
def test_outside_video_found_after_move_and_relink(tmp_path):
    vids = tmp_path / "videos"
    vids.mkdir()
    (vids / "v1.mp4").write_bytes(b"x")
    (vids / "v2.mp4").write_bytes(b"x")
    p = Project(name="e")
    p.save(tmp_path / "lab" / "e.mmaze")
    t1 = p.add_test(str(vids / "v1.mp4"), "A1")
    t2 = p.add_test(str(vids / "v2.mp4"), "A2")
    assert t1.video.startswith("..")
    p.save()
    d = json.loads((p.path / PROJECT_FILE).read_text(encoding="utf-8"))
    assert d["tests"][0]["video_abs"] == str((vids / "v1.mp4").resolve())
    moved = tmp_path / "elsewhere" / "deeper" / "e.mmaze"
    moved.parent.mkdir(parents=True)
    shutil.move(str(p.path), str(moved))
    q = Project.load(moved)
    assert os.path.exists(q.abs_path(q.tests[0].video)) and q.missing_videos() == []
    assert Project.from_dict(q.to_dict()).to_dict() == Project.from_dict(q.to_dict()).to_dict()
    # the videos move too: relink by name
    new = tmp_path / "archive" / "2024"
    new.mkdir(parents=True)
    for f in vids.iterdir():
        shutil.move(str(f), str(new / f.name.upper()))  # letter case ignored
    assert {t.id for t in q.missing_videos()} == {t1.id, t2.id}
    res = relink_videos(q, tmp_path / "archive")
    assert set(res["relinked"]) == {t1.id, t2.id} and not res["not_found"]
    assert q.missing_videos() == []
    # two files of the same name: the one under the matching folder wins, a tie is left alone
    (tmp_path / "x" / "a").mkdir(parents=True)
    (tmp_path / "x" / "b").mkdir(parents=True)
    (tmp_path / "x" / "a" / "s.mp4").write_bytes(b"x")
    (tmp_path / "x" / "b" / "s.mp4").write_bytes(b"x")
    t3 = q.add_test("", "A3")
    t3.video = "../gone/b/s.mp4"
    t4 = q.add_test("", "A4")
    t4.video = "../gone/c/s.mp4"
    res = relink_videos(q, tmp_path / "x", [t3, t4])
    assert res["relinked"][t3.id].endswith(os.path.join("b", "s.mp4")) and res["ambiguous"] == [t4.id]


# ---------------------------------------------------------------- lock file
def _dead_pid() -> int:
    proc = subprocess.Popen([sys.executable, "-c", "pass"])
    proc.wait()
    return proc.pid


def test_lock_file(tmp_path):
    d = tmp_path / "e.mmaze"
    d.mkdir()
    assert explock.acquire(d) is None and explock.is_mine(explock.read(d))
    assert explock.held_by_other(d) is None  # our own
    explock.release(d)
    assert not explock.lock_path(d).exists()
    # a live process of this computer (our parent)
    live = {**explock.me(), "pid": os.getppid()}
    explock.lock_path(d).write_text(json.dumps(live), encoding="utf-8")
    assert explock.held_by_other(d)["pid"] == os.getppid()
    assert explock.acquire(d)["pid"] == os.getppid()  # refused: nothing written
    assert explock.read(d)["pid"] == os.getppid()
    explock.release(d)  # not ours: left alone
    assert explock.lock_path(d).exists()
    assert "process" in explock.describe(live)
    # stale: the process is gone
    explock.lock_path(d).write_text(json.dumps({**live, "pid": _dead_pid()}), encoding="utf-8")
    assert explock.held_by_other(d) is None and explock.acquire(d) is None and explock.is_mine(explock.read(d))
    # another computer: cannot be checked, counts as live; open anyway takes it over
    explock.lock_path(d).write_text(json.dumps({**live, "host": "other-mac"}), encoding="utf-8")
    assert explock.held_by_other(d)["host"] == "other-mac"
    assert explock.acquire(d, force=True) is None and explock.is_mine(explock.read(d))
    explock.lock_path(d).write_text("garbage", encoding="utf-8")
    assert explock.held_by_other(d) is None


def test_read_only_project_refuses_to_save(tmp_path):
    p = _project(tmp_path)
    p.read_only = True
    with pytest.raises(ValueError, match="read-only"):
        p.save()
    with pytest.raises(ValueError):
        p.restore_backup(p.path / PROJECT_FILE)


def _side_file(p, test, animal=None, owner=None, test_id=None):
    path = autosave.path_for(p, test)
    d = {"version": 1, "meta": {"test_id": test.id if test_id is None else test_id,
                                "animal": test.animal_id if animal is None else animal, "apparatus": "OF",
                                "stage": "", "trial": 1},
         "fps": 10.0, "duration_s": 0.0,
         "cols": {**{c: [1.0, 2.0, 3.0] for c in COLUMNS}, "t": [0.0, 0.1, 0.2], "detected": [True] * 3},
         "events": [], "saved_at": "2099-01-01T00:00:00"}
    autosave.write(path, d)
    if owner is not None:  # rewrite with another owner
        dd = autosave.read(path)
        dd["owner"] = owner
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(dd, fh)
    return path


def test_recover_skips_running_owners_and_locked_experiments(tmp_path):
    p = Project(name="e")
    p.apparatus.append(Apparatus(name="OF", arena=rect(0, 0, 100, 100)))
    p.save(tmp_path / "e.mmaze")
    t = p.add_test("", "A1", "OF")
    p.save()
    live_owner = {"host": explock.me()["host"], "pid": os.getppid()}
    f = _side_file(p, t, owner=live_owner)
    assert autosave.recover(p) == [] and os.path.exists(f)  # its writer is still running
    # the experiment is locked by a live program: nothing is touched
    _side_file(p, t, owner={"host": explock.me()["host"], "pid": _dead_pid()})
    explock.lock_path(p.path).write_text(json.dumps({**explock.me(), "pid": os.getppid()}), encoding="utf-8")
    assert autosave.recover(p) == [] and os.path.exists(f)
    explock.lock_path(p.path).unlink()
    p.read_only = True
    assert autosave.recover(p) == []
    p.read_only = False
    rec = autosave.recover(p)  # the writer is gone: a crash
    assert [x.id for x in rec] == [t.id] and not os.path.exists(f)


def test_recover_checks_the_animal(tmp_path):
    p = Project(name="e")
    p.apparatus.append(Apparatus(name="OF", arena=rect(0, 0, 100, 100)))
    p.save(tmp_path / "e.mmaze")
    t = p.add_test("", "B7", "OF")  # the id of a deleted test of A1, given to another animal's test
    p.save()
    _side_file(p, t, animal="A1")
    rec = autosave.recover(p)
    assert len(rec) == 1 and rec[0].id != t.id and rec[0].animal_id == "A1" and t.status == "pending"


def test_project_json_has_no_nan(tmp_path):
    p = _project(tmp_path)
    p.variables["x"] = float("nan")
    p.tests[0].result_variables["y"] = float("inf")
    p.settings_extra["z"] = [1.0, float("-inf")]
    p.save()
    text = (p.path / PROJECT_FILE).read_text(encoding="utf-8")
    json.loads(text, parse_constant=lambda c: pytest.fail(f"non-standard JSON {c}"))
    q = Project.load(p.path)
    assert q.variables["x"] is None and q.settings_extra["z"] == [1.0, None]
    # files of older versions with NaN still load
    (p.path / PROJECT_FILE).write_text(text.replace('"x": null', '"x": NaN'), encoding="utf-8")
    assert math.isnan(Project.load(p.path).variables["x"])


# ---------------------------------------------------------------- Save As
def test_save_as_own_folder_other_spelling_keeps_tracks(tmp_path):
    p = _project(tmp_path, "Open field.mmaze")
    tracks = sorted(x.name for x in (p.path / "tracks").iterdir())
    assert tracks
    variants = [tmp_path / "." / "Open field.mmaze"]
    if os.path.exists(tmp_path / "OPEN FIELD.mmaze"):  # case-insensitive disk (macOS, Windows)
        variants.append(tmp_path / "open field.mmaze")
    link = tmp_path / "link"
    try:
        link.symlink_to(tmp_path, target_is_directory=True)
        variants.append(link / "Open field.mmaze")
    except OSError:
        pass
    for v in variants:
        assert same_folder(v, p.path)
        assert p.save_as(v) is False  # a plain save
        assert sorted(x.name for x in (p.path / "tracks").iterdir()) == tracks
        assert len(p.load_tracks(p.tests[0])) == 1


def test_save_as_replaces_an_existing_experiment(tmp_path):
    p = _project(tmp_path, "a.mmaze")
    other = _project(tmp_path / "sub", "b.mmaze", n_tests=3)
    stale = other.path / "tracks" / "test_0003.csv"
    assert stale.exists()
    assert p.save_as(other.path) is True
    assert p.path == other.path and not stale.exists()
    assert len(p.load_tracks(p.tests[0])) == 1
    assert not [x for x in other.path.iterdir() if x.name.startswith("tracks.")]
    assert (tmp_path / "a.mmaze" / "tracks" / "test_0001.csv").exists()  # the original is untouched


def test_save_as_failure_leaves_destination(tmp_path, monkeypatch):
    p = _project(tmp_path, "a.mmaze")
    other = _project(tmp_path / "sub", "b.mmaze", n_tests=3)
    old_path = p.path
    monkeypatch.setattr(Project, "save", lambda self, path=None: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        p.save_as(other.path)
    assert p.path == old_path
    assert (other.path / "tracks" / "test_0003.csv").exists()  # the old tracks are back
    assert not [x for x in other.path.iterdir() if x.name.startswith("tracks.")]


# ---------------------------------------------------------------- spreadsheets
def test_xlsx_heading_row_and_animals_csv_formula_protection(tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "r.xlsx"
    export.write_xlsx({"R": [{"=HYPERLINK(\"x\")": 1, "+cmd|' /C calc'!A0": "=1+1"}]}, out)
    ws = load_workbook(out).active
    assert ws.cell(1, 1).data_type == "s" and ws.cell(1, 2).data_type == "s" and ws.cell(2, 2).data_type == "s"
    assert not list(tmp_path.glob("*.tmp"))
    p = _project(tmp_path)
    p.animals[0].id = "=cmd()"
    p.animals[0].notes = "@SUM(A1)"
    p.animal_fields = ["=F"]
    f = tmp_path / "animals.csv"
    export.export_animals(p, f)
    text = f.read_text(encoding="utf-8")
    assert "'=cmd()" in text and "'@SUM(A1)" in text and "'=F" in text


def test_xlsx_write_is_atomic(tmp_path, monkeypatch):
    out = tmp_path / "r.xlsx"
    export.write_xlsx({"R": [{"a": 1}]}, out)
    good = out.read_bytes()
    from openpyxl import Workbook

    monkeypatch.setattr(Workbook, "save", lambda self, fh: (fh.write(b"half"), (_ for _ in ()).throw(OSError("x"))))
    with pytest.raises(OSError):
        export.write_xlsx({"R": [{"a": 2}]}, out)
    assert out.read_bytes() == good and not list(tmp_path.glob("*.tmp"))


def test_dbf_small_numbers_stay_numeric(tmp_path):
    rows = [{"v": 1e-05}, {"v": 123.5}, {"v": -2.5e-07}, {"v": 3}]
    typ, width, dec, enc = export._dbf_field([r["v"] for r in rows])
    assert typ == "N" and dec == 8
    f = tmp_path / "t.dbf"
    export.write_dbf(rows, f)
    head, body = export.read_dbf(f)
    assert [b[0] for b in body] == pytest.approx([1e-05, 123.5, -2.5e-07, 3])
    # too many digits to fit a numeric field: text
    assert export._dbf_field([1e25, 1.0])[0] == "C"


def test_raw_data_names_do_not_collide(tmp_path):
    p = _project(tmp_path)
    t = p.tests[0]
    t.extra_animals = ["A/1"]
    t.animal_id = "A_1"
    p.save_tracks(t, [_track(), _track()])
    files = export.export_raw_data(p, tmp_path / "raw", tests=[t], derived=False)
    assert len(files) == 2 and len({f.name for f in files}) == 2


def test_track_write_temp_names_are_unique(tmp_path):
    f = tmp_path / "x.txt"
    with atomic_write(f) as a, atomic_write(f) as b:  # two writers at once (threads, processes, computers)
        assert a.name != b.name
        a.write("a")
        b.write("b")
    assert f.read_text() in ("a", "b") and not list(tmp_path.glob("*.tmp"))


# ---------------------------------------------------------------- results errors and blind testing
def test_one_bad_track_does_not_lose_all_results(tmp_path):
    p = _project(tmp_path)
    p.track_path(p.tests[0]).write_text("# fps=10\nt,x,y\nnot,a,number\n", encoding="utf-8")
    rows = p.results()
    assert len(rows) == 2 and ERROR_COLUMN in rows[0] and ERROR_COLUMN not in rows[1]
    html = export.html_report(p, tmp_path / "r.html", stats_measures=["Total distance (px)"]).read_text("utf-8")
    assert "Problems" in html and "Test 1" in html


def test_blind_group_column_shows_codes(tmp_path):
    p = _project(tmp_path)
    p.blind = True
    codes = workflow.blind_codes(p)
    rows = p.results()
    assert {r["Group"] for r in rows} == {codes["Saline"], codes["Drug"]}
    sheets, _ = export.results_workbook(p, rows)
    assert {r["Group"] for r in sheets["Animals"]} == {codes["Saline"], codes["Drug"]}
    html = export.html_report(p, tmp_path / "r.html").read_text("utf-8")
    html = re.sub(r"data:image/png;base64,[A-Za-z0-9+/=]*", "", html)  # base64 plots can spell "Drug" by chance
    assert "Saline" not in html and "Drug" not in html and codes["Drug"] in html
    p.blind = False
    assert {r["Group"] for r in p.results()} == {"Saline", "Drug"}


# ---------------------------------------------------------------- ANY-maze XML
AM = """<?xml version="1.0" encoding="utf-8"?>
<Experiment><Title>T</Title>
  <Animal><Number>1</Number><Treatment>Drug</Treatment>
    <Test><Number>1</Number><Date>{d1}</Date><Time>10:00:00</Time><Apparatus>Box</Apparatus>
      <Zone><Name>Arena</Name><Centre><x>100</x><y>100</y></Centre>
        <BoundingBox><x>0</x><y>0</y><w>200</w><h>200</h></BoundingBox></Zone>
      <r><tm>0</tm><c><x>10</x><y>10</y></c></r><r><tm>0.04</tm><c><x>11</x><y>10</y></c></r>
    </Test>
    <Test><Number>2</Number><Date>{d2}</Date><Time>11:00:00</Time><Apparatus>Box</Apparatus>
      <r><tm>0</tm><c><x>10</x><y>10</y></c></r><r><tm>0.04</tm><c><x>11</x><y>10</y></c></r>
    </Test>
  </Animal>
</Experiment>
"""


def test_anymaze_dates_order_from_the_whole_file(tmp_path):
    assert day_first(["05/03/2024", "13/03/2024"]) and not day_first(["03/05/2024", "03/13/2024"])
    assert day_first(["05/03/2024"])  # nothing tells: day first
    f = tmp_path / "us.xml"
    f.write_text(AM.format(d1="03/05/2024", d2="03/13/2024"), encoding="utf-8")
    tests = read_anymaze_xml(f)["animals"][0]["tests"]
    assert [t["recorded_at"][:10] for t in tests] == ["2024-03-05", "2024-03-13"]
    f.write_text(AM.format(d1="05/03/2024", d2="13/03/2024"), encoding="utf-8")
    tests = read_anymaze_xml(f)["animals"][0]["tests"]
    assert [t["recorded_at"][:10] for t in tests] == ["2024-03-05", "2024-03-13"]


def test_anymaze_nameless_animals_do_not_merge(tmp_path):
    f = tmp_path / "a.xml"
    f.write_text(AM.format(d1="2024-03-05", d2="2024-03-06"), encoding="utf-8")
    p = Project(name="imp")
    p.save(tmp_path / "imp.mmaze")
    r1 = import_anymaze_xml(p, f)
    r2 = import_anymaze_xml(p, f)
    assert r1["animals"] == ["Animal 1"] and r2["animals"] == ["Animal 1 (2)"]
    assert len(p.get_animal("Animal 1 (2)").id) and {t.animal_id for t in r2["tests"]} == {"Animal 1 (2)"}


# ---------------------------------------------------------------- overlay video export
def test_overlay_video_export_safety(tmp_path, monkeypatch):
    from manymaze.core import videoexport
    from manymaze.core.demo import create_demo_project

    p = create_demo_project(tmp_path / "d.mmaze", n_per_group=1, seconds=2)
    t = next(x for x in p.tests if p.has_track(x))
    src = p.abs_path(t.video)
    size = os.path.getsize(src)
    with pytest.raises(ValueError):
        videoexport.export_video(p, t, src)
    assert os.path.getsize(src) == size
    out = tmp_path / "out.mp4"
    out.write_bytes(b"previous export")
    assert videoexport.export_video(p, t, out, should_stop=lambda: True) is None
    assert out.read_bytes() == b"previous export" and not list(tmp_path.glob("*.part*"))
    monkeypatch.setattr(videoexport.OverlayRenderer, "render", lambda *a: (_ for _ in ()).throw(RuntimeError("x")))
    with pytest.raises(RuntimeError):
        videoexport.export_video(p, t, out)
    assert out.read_bytes() == b"previous export" and not list(tmp_path.glob("*.part*"))
    monkeypatch.undo()
    assert videoexport.export_video(p, t, out) == out and out.stat().st_size > 1000
    assert not list(tmp_path.glob("*.part*"))


# ---------------------------------------------------------------- CLI
def test_cli_workers_env_and_wide_options(tmp_path, monkeypatch, capsys):
    p = _project(tmp_path)
    p.tests[0].video = "missing.mp4"
    p.save()
    monkeypatch.setenv("MANYMAZE_WORKERS", "four")
    from manymaze.core.batch import default_workers

    with pytest.raises(ValueError, match="MANYMAZE_WORKERS"):
        default_workers(3)
    with pytest.raises(SystemExit) as e:
        cli.main(["project", str(p.path), "results", "--wide", "-o", str(tmp_path / "x.xml")])
    assert "--wide" in str(e.value)
    with pytest.raises(SystemExit) as e:
        cli.main(["project", str(p.path), "events", "-o", str(tmp_path / "x.xml")])
    assert "XML" in str(e.value)
    out = tmp_path / "w.csv"
    m = next(c for c in p.results()[0] if c.startswith("Total distance"))
    cli.main(["project", str(p.path), "results", "--wide", "--column", m, "-o", str(out)])
    head = out.read_text(encoding="utf-8").splitlines()[0]
    assert m in head and "Mean speed" not in head
    with pytest.raises(SystemExit):
        cli.main(["project", str(p.path), "results", "--wide", "--column", "Nonsense", "-o", str(out)])


def test_cli_track_refuses_locked_and_reports_save_failure(tmp_path, monkeypatch):
    p = _project(tmp_path)
    explock.lock_path(p.path).write_text(json.dumps({**explock.me(), "pid": os.getppid()}), encoding="utf-8")
    with pytest.raises(SystemExit) as e:
        cli.main(["project", str(p.path), "track"])
    assert "open in" in str(e.value)
    explock.lock_path(p.path).unlink()
    import manymaze.core.batch as batch

    monkeypatch.setattr(batch, "track_tests", lambda *a, **k: {"tracked": [1], "errors": [], "workers": 1,
                                                               "cancelled": False})
    monkeypatch.setattr(Project, "save", lambda self, path=None: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(SystemExit) as e:
        cli.main(["project", str(p.path), "track"])
    assert "tracks were written (tests 1)" in str(e.value) and "disk full" in str(e.value)
    assert not explock.lock_path(p.path).exists()  # released


def test_cli_relink(tmp_path, capsys):
    p = Project(name="e")
    p.save(tmp_path / "e.mmaze")
    t = p.add_test("", "A1")
    t.video = "../old/v.mp4"
    p.save()
    (tmp_path / "new").mkdir()
    (tmp_path / "new" / "v.mp4").write_bytes(b"x")
    cli.main(["project", str(p.path), "relink", "--folder", str(tmp_path / "new")])
    assert "Relinked 1" in capsys.readouterr().out
    assert Project.load(p.path).missing_videos() == []
