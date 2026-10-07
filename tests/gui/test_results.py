"""Smoke tests for the Results page."""

import csv
import shutil

import pytest
from PySide6.QtCore import QElapsedTimer, Qt
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from manymaze.core.demo import create_demo_project
from manymaze.core.project import Project
from manymaze.gui.main_window import MainWindow
from manymaze.gui.pages import _results_cache
from manymaze.gui.pages.results import measure_category, _names

app = QApplication.instance() or QApplication([])


@pytest.fixture(scope="module")
def demo_dir(tmp_path_factory):
    d = tmp_path_factory.mktemp("demo") / "d.mmaze"
    p = create_demo_project(d, n_per_group=3, seconds=8)
    p.analysis.bin_length_s = 4.0
    p.save()
    return d


@pytest.fixture
def page(demo_dir, tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: QMessageBox.Ok)
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda *a, **k: True)
    d = tmp_path / "d.mmaze"
    shutil.copytree(demo_dir, d)
    w = MainWindow()
    w.set_project(Project.load(d))
    p = w.goto("ResultsPage")
    p.wait_loaded()
    yield p
    w.dirty = False
    w.close()


def wait_workers(page, ms=120000):
    t = QElapsedTimer()
    t.start()
    while getattr(page, "_workers", None) and t.elapsed() < ms:
        for w in list(page._workers):
            w.wait(50)
        app.processEvents()
    app.processEvents()


def test_table_populated(page):
    cols = page.model.columns
    assert page.model.rowCount() == 6
    for c in ("Test", "Animal", "Group", "Stage", "Trial", "Total distance (cm)", "Centre: time (%)",
              "Rearing: count", "Defecation: count"):
        assert c in cols
    assert "Period" not in cols  # whole-test view
    assert {r["Group"] for r in page.rows} == {"Control", "Anxious"}
    # numeric formatting: 3 decimals, NaN blank
    ci = cols.index("Centre: time (%)")
    txt = page.proxy.index(0, ci).data()
    assert txt == "" or len(txt.split(".")[-1]) == 3
    # sorting by a numeric column
    page.table.sortByColumn(ci, Qt.DescendingOrder)
    vals = [page.shown_rows()[i]["Centre: time (%)"] for i in range(6)]
    assert vals == sorted(vals, reverse=True)
    # detail plots for the selected row
    page.table.selectRow(0)
    page.render_all()
    assert page.track_canvas.figure.axes and page.heat_canvas.figure.axes and page.speed_canvas.figure.axes
    # filters
    page.group_combo.setCurrentIndex(page.group_combo.findData("Control"))
    assert page.proxy.rowCount() == 3
    page.group_combo.setCurrentIndex(0)
    assert page.proxy.rowCount() == 6


def test_categories(page):
    names = _names(page.project)
    assert measure_category("Centre: time (%)", names) == ("Zones", "Centre")
    assert measure_category("Rearing: count", names) == ("Behaviours", "Rearing")
    assert measure_category("Total distance (cm)", names) == ("General", "")
    assert measure_category("Escape latency (s)", names) == ("Test-specific", "")
    cats = [page.tree.topLevelItem(i).text(0) for i in range(page.tree.topLevelItemCount())]
    assert cats[:2] == ["Information", "General"] and "Zones" in cats and "Behaviours" in cats
    info = page.tree.topLevelItem(0)
    app_item = next(info.child(i) for i in range(info.childCount()) if info.child(i).text(0) == "Apparatus")
    assert "Apparatus" in page.model.columns
    app_item.setCheckState(0, Qt.Unchecked)
    app.processEvents()
    assert "Apparatus" not in page.model.columns and "Animal" in page.model.columns


def test_column_chooser(page):
    page.set_visible_measures(["Total distance (cm)", "Centre: time (%)"])
    assert page.visible_measures() == ["Total distance (cm)", "Centre: time (%)"]
    assert page.model.columns[-2:] == ["Total distance (cm)", "Centre: time (%)"]
    assert "Rearing: count" not in page.model.columns
    # check an item in the tree → column appears
    zones = next(page.tree.topLevelItem(i) for i in range(page.tree.topLevelItemCount())
                 if page.tree.topLevelItem(i).text(0) == "Behaviours")
    rearing = next(zones.child(i) for i in range(zones.childCount()) if zones.child(i).text(0) == "Rearing")
    count = next(rearing.child(i) for i in range(rearing.childCount()) if rearing.child(i).text(0) == "count")
    count.setCheckState(0, Qt.Checked)
    app.processEvents()
    assert "Rearing: count" in page.model.columns
    # unchecking a parent hides all its children
    rearing.setCheckState(0, Qt.Unchecked)
    app.processEvents()
    assert "Rearing: count" not in page.model.columns
    # search + select all
    page.search.setText("Corner 1")
    page.set_all_measures(True)
    assert "Corner 1: entries" in page.model.columns and "Corner 2: entries" not in page.model.columns
    page.search.setText("")
    page.set_all_measures(True)
    assert len(page.visible_measures()) == len(page.measure_columns())
    # selection kept in memory on the project
    page.set_visible_measures(["Centre: time (%)"])
    assert "Total distance (cm)" in page.hidden


def test_exports(page, tmp_path, monkeypatch):
    page.set_visible_measures(["Total distance (cm)", "Centre: time (%)"])
    out = tmp_path / "r.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    page.save_table()
    with open(out) as f:
        rows = list(csv.reader(f))
    assert rows[0] == page.shown_columns()
    assert rows[0][-2:] == ["Total distance (cm)", "Centre: time (%)"]
    assert len(rows) == 7
    # clipboard (TSV)
    page.table.clearSelection()
    page.copy_to_clipboard()
    lines = QApplication.clipboard().text().splitlines()
    assert lines[0].split("\t") == page.shown_columns() and len(lines) == 7
    # time periods + xlsx
    page.seg_check.setChecked(True)
    page.wait_loaded()
    assert page.segmented and "Period" in page.model.columns
    assert {r["Period"] for r in page.rows} == {"Whole test", "0-4 s", "4-8 s"}
    xo = tmp_path / "r.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(xo), ""))
    page.save_table(suffix=".xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(xo)
    assert "Results" in wb.sheetnames and "Time periods" in wb.sheetnames
    head = [c.value for c in wb["Results"][1]]
    assert head[-2:] == ["Total distance (cm)", "Centre: time (%)"] and "Rearing: count" not in head
    assert wb["Results"].max_row == 7
    tp = [c.value for c in wb["Time periods"][1]]
    assert "Period" in tp and tp[-1] == "Centre: time (%)"
    assert wb["Time periods"].max_row == 13
    assert {"Zone visits", "Animals", "Tests", "Settings"} <= set(wb.sheetnames)
    assert any(r[1].value.startswith("mANY-MAZE") for r in wb["Settings"].iter_rows(min_row=2))
    # period filter
    page.period_combo.setCurrentIndex(page.period_combo.findData("0-4 s"))
    assert page.proxy.rowCount() == 6


def test_figures_and_report(page, tmp_path):
    page.table.selectRow(1)
    page.render_all()
    png = page.save_figure(str(tmp_path / "track.svg"))
    assert png and (tmp_path / "track.svg").exists()
    page.group_heatmaps()
    wait_workers(page)
    fig = page.groups_canvas.figure
    assert len([a for a in fig.axes if a.images]) == 2
    page.save_figure(str(tmp_path / "groups.png"))
    assert (tmp_path / "groups.png").stat().st_size > 1000
    rep = tmp_path / "report.html"
    w = page.html_report(str(rep), stats_measures=["Centre: time (%)"], include_plots=False)
    assert w.wait(120000)
    app.processEvents()
    text = rep.read_text()
    assert "Statistics" in text and "Centre: time (%)" in text


def test_cache_shared(page):
    p = page.project
    rows = _results_cache.cached_rows(p, False)
    assert rows is page.rows
    p.analysis.min_freeze_s = 2.0  # analysis change invalidates the cache
    assert _results_cache.cached_rows(p, False) is None
    page.reload()
    page.wait_loaded()
    assert _results_cache.cached_rows(p, False) is page.rows
    # so do the I/O log, saved procedure variables and I/O devices
    for change in (lambda: p.tests[0].io_events.append({"t": 1.0, "device": "d", "channel": 1, "value": 1}),
                   lambda: p.tests[0].result_variables.update(score=1.0), lambda: p.io_devices.append({})):
        fp = _results_cache.fingerprint(p)
        change()
        assert _results_cache.fingerprint(p) != fp
