"""Results page: charts tab, plot options, behaviour / aligned heat maps, selection exports, XML / raw data and
video export with overlays."""

import csv
import shutil
import time

import pytest
from PySide6.QtCore import QItemSelection, QItemSelectionModel
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QMessageBox

from manymaze.core.demo import create_demo_project
from manymaze.core.export import read_experiment_xml
from manymaze.core.project import Project
from manymaze.gui.main_window import MainWindow
from manymaze.gui.pages import _results_cache

app = QApplication.instance() or QApplication([])


@pytest.fixture(scope="module")
def demo_dir(tmp_path_factory):
    d = tmp_path_factory.mktemp("demo") / "d.mmaze"
    p = create_demo_project(d, n_per_group=1, seconds=6)
    p.analysis.bin_length_s = 3.0
    p.save()
    return d


@pytest.fixture
def page(demo_dir, tmp_path, monkeypatch):
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: QMessageBox.Ok)
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda *a, **k: True)
    d = tmp_path / "d.mmaze"
    shutil.copytree(demo_dir, d)
    w = MainWindow()
    w.set_project(Project.load(d))
    p = w.goto("ResultsPage")
    p.wait_loaded()
    yield p
    w.dirty = False
    w.close()


def wait_workers(page, ms=120000):
    t0 = time.time()
    while getattr(page, "_workers", None) and time.time() - t0 < ms / 1000:
        app.processEvents()
        time.sleep(0.03)
    app.processEvents()


def test_charts_tab(page, tmp_path):
    page.main_tabs.setCurrentIndex(1)
    ch = page.charts
    assert ch.test_combo.count() == 2 and ch.track is not None
    n_params = sum(1 for i in range(ch.param_list.count()) if ch.param_list.item(i).data(0x0100))
    assert n_params >= 40
    assert ch.checked_params() == ["Speed", "Freezing"]
    ch.set_selection(test_id=page.project.tests[1].id, params=["Speed", "Distance from wall", "Centre: in zone"],
                     bands=["Centre"])
    assert ch.test.id == page.project.tests[1].id
    axes = ch.figure.axes
    assert len(axes) == 4  # 3 parameters + scored events strip
    assert axes[0].get_legend() is not None
    rows = ch.measure(1.0, 4.0)
    # charted in the order of the parameter list
    assert [r["parameter"] for r in rows] == ["Distance from wall", "Speed", "Centre: in zone"]
    assert ch.meas_table.rowCount() == 3 and rows[0]["duration"] == pytest.approx(3.0)
    pk = ch.find_peaks()
    assert set(pk) == {"Speed", "Distance from wall"}
    ch.measure_check.setChecked(True)
    assert len(ch._spans) == len(ch.figure.axes)
    # period zoom
    ch.period_combo.setCurrentIndex(ch.period_combo.findText("3-6 s"))
    ch.redraw()
    assert ch.figure.axes[0].get_xlim() == pytest.approx((3.0, 6.0))
    assert ch.save_image(str(tmp_path / "chart.png")) and (tmp_path / "chart.png").stat().st_size > 2000
    ch.copy_image()
    assert not QApplication.clipboard().image().isNull()
    out = ch.export_data(str(tmp_path / "chart.tsv"))
    head = open(out).readline().rstrip("\n").split("\t")
    assert head == ["Time (s)", "Distance from wall (cm)", "Speed (cm/s)", "Centre: in zone"]
    # context menu action jumps from the table to the chart of the selected test
    page.main_tabs.setCurrentIndex(0)
    page.table.selectRow(0)
    page.show_charts()
    assert page.main_tabs.currentIndex() == 1 and ch.test.id == page.current_row()["Test"]


def test_track_and_heatmap_options(page):
    page.table.selectRow(0)
    page.set_plot_options(color_by="speed", markers=True, part="head")
    page.render_all()
    fig = page.track_canvas.figure
    assert len(fig.axes) == 2 and "speed" in fig.axes[1].get_ylabel()
    assert page.color_combo.findData("Distance from wall") >= 0
    page.set_plot_options(color_by="Distance from wall")
    page.render_all()
    assert "wall" in page.track_canvas.figure.axes[1].get_ylabel()
    page.set_plot_options(split=True, color_by="time")
    page.render_all()
    assert len([a for a in page.track_canvas.figure.axes if a.get_title()]) == 2  # 0-3 s, 3-6 s
    page.set_plot_options(split=False)
    # behaviour heat map with % normalisation, then a fixed scale
    assert page.heat_of.findData("Freezing") >= 0 and page.heat_of.findData("Rearing: active") >= 0
    page.tabs.setCurrentIndex(1)
    page.set_plot_options(heat_of="Rearing: active", norm="percent")
    ax = page.heat_canvas.figure.axes
    assert ax[0].get_title() == "Rearing: active" and "% of time" in ax[1].get_ylabel()
    page.set_plot_options(norm="fixed", vmax=0.5)
    assert page.heat_max.isEnabled()
    assert page.heat_canvas.figure.axes[0].images[0].get_clim() == (0, 0.5)


def test_group_heatmaps_alignment(page):
    p = page.project
    page.table.selectRow(1)
    rows_before = _results_cache.cached_rows(p, False)
    page.set_plot_options(align="rot90")
    t = p.get_test(page.current_row()["Test"])
    assert t.variables["heatmap_transform"] == "rot90" and page.main.dirty
    assert _results_cache.cached_rows(p, False) is rows_before  # display-only: results stay cached
    page.set_plot_options(heat_of="Freezing", norm="relative")
    page.group_heatmaps()
    wait_workers(page)
    fig = page.groups_canvas.figure
    assert len([a for a in fig.axes if a.images]) == 2
    assert "relative" in fig.axes[-1].get_ylabel()
    page.set_plot_options(align="none")
    assert "heatmap_transform" not in t.variables


def test_selection_and_experiment_exports(page, tmp_path, monkeypatch):
    page.set_visible_measures(["Total distance (cm)", "Centre: time (%)", "Thigmotaxis (%)"])
    cols = page.model.columns
    # select a 2 × 2 cell range
    sm = page.table.selectionModel()
    page.table.clearSelection()
    c0 = cols.index("Total distance (cm)")
    c1 = c0 + 1
    sel = QItemSelection(page.proxy.index(0, c0), page.proxy.index(1, c1))
    sm.select(sel, QItemSelectionModel.ClearAndSelect)
    rows, rcols = page.selection_range()
    assert rcols == cols[c0:c1 + 1] and len(rows) == 2
    lines = page.clipboard_text().splitlines()
    assert lines[0] == "\t".join(cols[c0:c1 + 1]) and len(lines) == 3
    page.copy_to_clipboard()
    assert QApplication.clipboard().text().splitlines()[0] == lines[0]
    out = page.save_table(str(tmp_path / "sel.tsv"), selection=True)
    assert open(out).read().splitlines()[0] == lines[0]
    page.save_table(str(tmp_path / "sel.xlsx"), selection=True)
    from openpyxl import load_workbook

    assert load_workbook(tmp_path / "sel.xlsx").active.max_row == 3
    # whole table as tab-separated text
    page.table.clearSelection()
    out = page.save_table(str(tmp_path / "all.tsv"))
    with open(out) as f:
        data = list(csv.reader(f, delimiter="\t"))
    assert data[0] == page.shown_columns() and len(data) == 3
    # experiment XML (background worker) and raw data per test
    page.export_xml(str(tmp_path / "exp.xml"))
    wait_workers(page)
    x = read_experiment_xml(tmp_path / "exp.xml")
    assert len(x["tests"]) == 2 and x["tests"][0]["tracks"] and x["tests"][0]["results"]
    page.export_raw(str(tmp_path / "raw"))
    wait_workers(page)
    assert len(list((tmp_path / "raw").glob("test_*.csv"))) == 2
    # report with heat map normalisation and charts
    page.html_report(str(tmp_path / "r.html"), stats_measures=["Centre: time (%)"], include_plots=True,
                     heatmap_norm="percent", chart_parameters=["Speed"])
    wait_workers(page)
    assert (tmp_path / "r.html").read_text().count("<img") >= 7


def test_video_export(page, tmp_path, monkeypatch):
    from manymaze.core.video import VideoSource
    from manymaze.gui.pages.results import VideoExportDialog

    page.table.selectRow(0)
    # dialog defaults → options
    dlg = VideoExportDialog(page)
    o = dlg.options()
    assert o.trail_s == 5.0 and o.speed == 1.0 and o.zones
    dlg.speed.setCurrentIndex(dlg.speed.findData(2.0))
    dlg.scale.setCurrentIndex(dlg.scale.findData(0.5))
    monkeypatch.setattr(VideoExportDialog, "exec", lambda self: QDialog.Accepted)
    monkeypatch.setattr(VideoExportDialog, "options", lambda self: o.__class__(speed=2.0, scale=0.5))
    out = tmp_path / "v.mp4"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    page.export_video()
    wait_workers(page)
    with VideoSource(str(out)) as v:
        assert v.width == 200 and 70 <= v.frame_count <= 80
